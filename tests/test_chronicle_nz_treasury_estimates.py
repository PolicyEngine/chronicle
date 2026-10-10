"""Publisher-cell fidelity for Treasury's Budget 2026 Social Development rows."""

from __future__ import annotations

import hashlib
import re
from collections import Counter
from copy import deepcopy
from dataclasses import replace
from decimal import Decimal
from functools import lru_cache
from pathlib import Path

import openpyxl
import pytest
import yaml
from pypdf import PdfReader

from chronicle.artifacts import build_r2_key
from chronicle.bundle import _dimension_label_reports, build_bundle_coverage
from chronicle.consumer_contract import (
    consumer_fact_rows,
    validate_consumer_fact_contract,
)
from chronicle.core import validate_facts
from chronicle.source_package import (
    SOURCE_PACKAGE_ALIASES,
    load_source_package,
    validate_source_package,
)
from chronicle.sources import build_source_cell_key, validate_source_cells
from chronicle.sources.rows import build_source_row_key, validate_source_rows
from chronicle.sources.specs import resolve_source_record

REPO_ROOT = Path(__file__).resolve().parents[1]
ALIAS = "treasury-estimates-of-appropriations-2026-27"
DIRECTORY = Path("treasury/estimates_of_appropriations_2026_27")
PACKAGE = REPO_ROOT / "packages" / DIRECTORY
DATA = REPO_ROOT / "db" / "data" / DIRECTORY
WORKBOOK = "b26-expenditure-data.xlsx"
PDF = "est26-v9-socdev.pdf"
WORKBOOK_SHA256 = "3fc6bba178c78c4a4b259c920a6f55307ec95a547353f340086c86fc2a26f5a0"
PDF_SHA256 = "85ff85e5b93a8b3e3ef7442a034eb8d3abdd0f1675f76646afad8311d8e30c23"
SOURCE_PAGE = "https://www.budget.govt.nz/budget/2026/estimates/data.htm"
WORKBOOK_URL = "https://www.budget.govt.nz/budget/excel/data/" + WORKBOOK
PDF_URL = "https://budget.govt.nz/budget/pdfs/estimates/v9/" + PDF
HEADERS = (
    "Department",
    "Vote",
    "App ID",
    "Parent ID",
    "Appropriation Name",
    "Category Name",
    "Group Type",
    "Appropriation or Category Type",
    "Restriction Type",
    "Functional Classification",
    "Amount $000",
    "Year",
    "Amount Type",
    "Periodicity",
    "Current Scope",
    "M Number",
    "Portfolio Name",
)
YEAR_COUNTS = {2022: 19, 2023: 18, 2024: 18, 2025: 19, 2026: 19, 2027: 19}
# Physical worksheet addresses, not the parser's compact virtual addresses.
EXACT_CELLS = (
    (4783, 12381, "Accommodation Assistance", 2025, "Actuals", 2_232_026),
    (4797, 12381, "Accommodation Assistance", 2026, "Estimated Actual", 2_308_335),
    (4816, 12381, "Accommodation Assistance", 2027, "Main Estimates", 2_322_160),
    (4821, 5406, "New Zealand Superannuation", 2027, "Main Estimates", 26_481_340),
    (
        4807,
        10803,
        "Jobseeker Support and Emergency Benefit",
        2027,
        "Main Estimates",
        5_018_386,
    ),
    (4808, 10804, "Sole Parent Support", 2027, "Main Estimates", 2_473_973),
    (4809, 10805, "Supported Living Payment", 2027, "Main Estimates", 3_023_208),
)
SCOPE = (
    "This appropriation is limited to payments for accommodation costs, paid in "
    "accordance with criteria set out in, or in delegated legislation made under, "
    "the Social Security Act 2018."
)
ACCOMMODATION_SUPPLEMENT_REFERENCE = (
    "Accommodation Supplement is paid under section 65 of the Social Security Act 2018"
)
AWAY_FROM_HOME_REFERENCE = (
    "Away from Home Allowance is paid under the Away from Home Allowance Welfare "
    "Programme pursuant to section 124(1)(d) of the Social Security Act 1964 saved "
    "by clause 21 of Schedule 1 of the Social Security Act 2018 as if it were a "
    "special assistance programme approved and established under section 101 of "
    "the Social Security Act 2018"
)


@lru_cache
def _package():
    return load_source_package(PACKAGE)


@lru_cache
def _payload():
    return yaml.safe_load((PACKAGE / "source_package.yaml").read_text())


@lru_cache
def _rows():
    return _package().build_source_rows(2026)


@lru_cache
def _cells():
    return _package().build_source_cells(2026, source_rows=_rows())


@lru_cache
def _facts():
    return _package().build_facts(2026, cells=_cells(), source_rows=_rows())


@lru_cache
def _records():
    return _package().build_source_records(2026, cells=_cells(), source_rows=_rows())


@lru_cache
def _specs():
    return _package().build_source_record_specs(2026)


@lru_cache
def _publisher_rows():
    """Read physical cells independently of Chronicle's row/cell compaction."""
    workbook = openpyxl.load_workbook(DATA / WORKBOOK, read_only=True, data_only=True)
    try:
        sheet = workbook["Raw Data"]
        assert (sheet.max_row, sheet.max_column) == (6452, 17)
        rows = tuple(sheet.iter_rows(values_only=True))
        assert rows[0] == HEADERS
        return {number: row for number, row in enumerate(rows[1:], start=2)}
    finally:
        workbook.close()


def _compact_address(physical_row, column="K"):
    cell = next(
        cell
        for cell in _cells()
        if cell.note == f"source_line_number={physical_row}"
        and cell.address.startswith(column)
    )
    return cell.address


def test_treasury_estimates_package_validation_and_alias():
    assert SOURCE_PACKAGE_ALIASES[ALIAS] == DIRECTORY
    assert _package().package_id == ALIAS
    report = validate_source_package(PACKAGE, year=2026)
    assert report.valid, report.to_dict()
    assert not report.warnings
    assert report.counts == {
        "record_set_count": 6,
        "row_count": 112,
        "measure_count": 6,
        "source_record_count": 112,
        "source_region_count": 6,
    }


@pytest.mark.parametrize(
    "entry,filename,sha256,size,url,fetched_at",
    (
        (
            2026,
            WORKBOOK,
            WORKBOOK_SHA256,
            918_842,
            WORKBOOK_URL,
            "2026-10-10T08:33:39Z",
        ),
        (
            "scope_evidence",
            PDF,
            PDF_SHA256,
            2_212_387,
            PDF_URL,
            "2026-10-10T08:31:09Z",
        ),
    ),
)
def test_treasury_estimates_immutable_artifacts_and_receipt_provenance(
    entry, filename, sha256, size, url, fetched_at
):
    manifest = yaml.safe_load((DATA / "manifest.yaml").read_text())
    assert manifest["source_id"] == "treasury"
    assert manifest["package_id"] == ALIAS
    assert manifest["source_page"] == SOURCE_PAGE
    assert manifest["published_at"] == "2026-05-28"
    assert manifest["license"] == "CC BY 4.0"
    assert manifest["license_url"] == "https://www.budget.govt.nz/about/copyright.htm"
    assert set(manifest["files"]) == {2026, "scope_evidence"}
    artifact = manifest["files"][entry]
    assert artifact["filename"] == filename
    assert artifact["source_url"] == url
    assert artifact["fetched_at"] == fetched_at
    assert artifact["sha256"] == sha256
    assert hashlib.sha256((DATA / filename).read_bytes()).hexdigest() == sha256
    assert (DATA / filename).stat().st_size == artifact["size_bytes"] == size
    key = build_r2_key(
        source_id="treasury",
        package_id=ALIAS,
        year=entry,
        sha256=sha256,
        filename=filename,
        prefix="raw/nz",
        package_path=PACKAGE,
    )
    assert key == f"raw/nz/treasury/{ALIAS}/{entry}/{sha256}/{filename}"
    assert artifact["storage"]["r2"] == {
        "provider": "r2",
        "bucket": "ledger-raw",
        "key": key,
        "uri": f"r2://ledger-raw/{key}",
    }


@pytest.mark.parametrize("row,app_id,component,year,status,raw", EXACT_CELLS)
def test_treasury_estimates_exact_published_amount_cells(
    row, app_id, component, year, status, raw
):
    publisher = _publisher_rows()[row]
    assert publisher[1] == "Social Development"
    assert publisher[2] == app_id
    assert publisher[4] == component
    assert publisher[7] == "Benefits or Related Expenses"
    assert publisher[10:13] == (raw, year, status)
    address = _compact_address(row)
    cells = {cell.address: cell for cell in _cells()}
    cell = cells[address]
    assert cell.raw_value == raw
    assert cell.note == f"source_line_number={row}"
    source_row = next(
        source_row for source_row in _rows() if source_row.row_number == row
    )
    assert cell.source_row_key == build_source_row_key(source_row)
    record = next(
        record for record in _records() if record.spec.selector.address == address
    )
    fact = next(
        fact for fact in _facts() if fact.source_record_id == record.source_record_id
    )
    assert fact.value == raw * 1000
    assert fact.measure.unit == "nzd"
    assert fact.filters == {
        "benefit_expense_component": component,
        "amount_type": status,
    }
    assert (fact.period.type, fact.period.value) == ("fiscal_year", year - 1)
    assert build_source_cell_key(cell) in fact.source_cell_keys
    assert cell.source_row_key in fact.source_row_keys


def test_treasury_estimates_complete_scope_and_compact_lineage():
    rows = _rows()
    assert len(rows) == 6451
    assert [row.row_number for row in rows] == list(range(2, 6453))
    assert all(tuple(row.values) == HEADERS for row in rows)
    assert {row.sheet_name for row in rows} == {"Raw Data"}
    assert validate_source_rows(rows).valid
    publisher_rows = _publisher_rows()
    selected = {
        row_number: row
        for row_number, row in publisher_rows.items()
        if row[1] == "Social Development" and row[7] == "Benefits or Related Expenses"
    }
    assert list(selected) == list(range(4713, 4825))
    assert Counter(row[11] for row in selected.values()) == YEAR_COUNTS
    assert len(_payload()["artifact"]["selected_rows"]) == 112
    assert _payload()["artifact"]["parser"] == "xlsx_table_full_rows"
    assert _payload()["artifact"]["sheet_name"] == "Raw Data"
    cells = _cells()
    assert len(cells) == 17 * 113
    assert {cell.sheet_name for cell in cells} == {"Raw Data"}
    assert validate_source_cells(cells).valid
    assert [cell.raw_value for cell in cells if cell.row_number == 1] == list(HEADERS)
    selected_keys = {
        build_source_row_key(row) for row in rows if row.row_number in selected
    }
    assert {
        cell.source_row_key for cell in cells if cell.row_number > 1
    } == selected_keys
    assert {cell.note for cell in cells if cell.row_number > 1} == {
        f"source_line_number={row}" for row in selected
    }
    assert {key for fact in _facts() for key in fact.source_row_keys} == selected_keys
    assert len(_facts()) == len(selected)


def test_treasury_estimates_every_fact_is_one_scaled_publisher_cell():
    cells = {cell.address: cell for cell in _cells()}
    records = {record.source_record_id: record for record in _records()}
    row_keys = {build_source_row_key(row) for row in _rows()}
    cell_keys = {build_source_cell_key(cell) for cell in _cells()}
    facts = _facts()
    assert validate_facts(facts).valid
    assert validate_consumer_fact_contract(facts).valid
    assert {fact.provenance_class for fact in facts} == {"administrative"}
    assert {fact.measure.unit for fact in facts} == {"nzd"}
    assert {fact.source.source_sha256 for fact in facts} == {WORKBOOK_SHA256}
    assert {fact.source.source_size_bytes for fact in facts} == {918_842}
    assert {fact.source.url for fact in facts} == {WORKBOOK_URL}
    assert {fact.source.raw_r2_uri for fact in facts} == {
        f"r2://ledger-raw/raw/nz/treasury/{ALIAS}/2026/{WORKBOOK_SHA256}/{WORKBOOK}"
    }
    for fact in facts:
        record = records[fact.source_record_id]
        assert record.spec.selector.address.startswith("K")
        assert record.spec.selector.end_address is None
        assert record.spec.divisor_selector is None
        assert record.spec.round_to is None
        assert record.spec.value_scale == 1000
        value_cell = cells[record.spec.selector.address]
        assert Decimal(str(fact.value)) == Decimal(str(value_cell.raw_value)) * 1000
        assert build_source_cell_key(value_cell) in fact.source_cell_keys
        assert fact.source_cell_keys and set(fact.source_cell_keys) <= cell_keys
        assert len(fact.source_row_keys) == 1
        assert set(fact.source_row_keys) <= row_keys
        assert fact.layout.table_record_kind == "detail"
        assert "growth" not in fact.measure.concept
        assert "average" not in fact.measure.concept


def test_treasury_estimates_publisher_period_status_and_labels():
    cells = {cell.address: cell for cell in _cells()}
    records = {record.source_record_id: record for record in _records()}
    assert Counter(fact.assertion for fact in _facts()) == {
        "observation": 74,
        "source_projection": 38,
    }
    assert Counter(fact.filters["amount_type"] for fact in _facts()) == {
        "Actuals": 74,
        "Estimated Actual": 19,
        "Main Estimates": 19,
    }
    for fact in _facts():
        record = records[fact.source_record_id]
        virtual_row = cells[record.spec.selector.address].row_number
        component = cells[f"E{virtual_row}"].raw_value
        year = cells[f"L{virtual_row}"].raw_value
        status = cells[f"M{virtual_row}"].raw_value
        assert fact.assertion == (
            "observation" if status == "Actuals" else "source_projection"
        )
        assert (fact.period.type, fact.period.value) == ("fiscal_year", year - 1)
        assert (fact.geography.level, fact.geography.id) == ("country", "NZ")
        assert fact.geography.name == "New Zealand"
        assert fact.entity.name == "government"
        coverage = fact.period_coverage
        assert coverage.basis == "fiscal"
        assert coverage.source_period_label == str(year)
        assert coverage.start_date == f"{year - 1}-07-01"
        assert coverage.end_date == f"{year}-06-30"
        assert fact.filters["benefit_expense_component"] == component
        assert fact.layout.groupby_value_label == component
        assert fact.layout.groupby_dimension == "treasury.appropriation"
        assert fact.layout.groupby_dimension_label == "Appropriation"
        assert (
            fact.dimension_value_labels["benefit_expense_component"][component]
            == component
        )
        assert fact.dimension_value_labels["amount_type"][status] == status
    assert _dimension_label_reports(ALIAS, consumer_fact_rows(_facts())) == ([], [])


def test_treasury_estimates_single_source_coverage():
    coverage = build_bundle_coverage(consumer_fact_rows(_facts()))
    assert coverage["fact_count"] == 112
    assert coverage["counts"]["by_source"] == {"treasury": 112}
    assert coverage["counts"]["by_geography"] == {"country:NZ": 112}
    assert coverage["counts"]["by_period"] == {
        f"fiscal_year:{year - 1}": count for year, count in YEAR_COUNTS.items()
    }
    assert coverage["counts"]["by_entity"] == {"government": 112}
    assert coverage["duplicates"] == {
        "aggregate_fact_keys": [],
        "semantic_fact_keys": [],
    }


def test_treasury_estimates_accommodation_assistance_msd_reference_period():
    # Compare the two independently asserted source cells and their declared
    # reference periods. No value is reconciled or emitted from their equality.
    fact = next(
        fact
        for fact in _facts()
        if fact.filters["benefit_expense_component"] == "Accommodation Assistance"
        and fact.period.value == 2024
    )
    msd = load_source_package("msd-annual-report-benefit-expenses-2025")
    msd_cells = msd.build_source_cells(2025)
    msd_value = next(cell.raw_value for cell in msd_cells if cell.address == "E4963")
    msd_record = next(
        spec
        for spec in msd.build_source_record_specs(2025)
        if spec.selector.address == "E4963"
    )
    assert msd_value == _publisher_rows()[4783][10] == 2_232_026
    assert fact.value == msd_value * 1000
    assert (
        (fact.period.type, fact.period.value)
        == (
            msd_record.period_type,
            msd_record.period,
        )
        == ("fiscal_year", 2024)
    )
    assert (
        fact.period_coverage.source_period_label
        == msd_record.period_coverage.source_period_label
        == "2025"
    )


def test_treasury_estimates_pdf_scope_evidence_never_emits_a_fact():
    reader = PdfReader(DATA / PDF)
    assert len(reader.pages) == 132
    # Printed pp.187, 243, 245 are zero-based PDF indices 4, 60, 62.
    summary = " ".join(reader.pages[4].extract_text().split())
    scope = " ".join(reader.pages[60].extract_text().split())
    conditions = " ".join(reader.pages[62].extract_text().split())
    assert "Accommodation Assistance (M63) (A25)" in summary
    assert "2,361,935 2,308,335 2,322,160" in summary
    assert SCOPE in summary and SCOPE in scope
    assert ACCOMMODATION_SUPPLEMENT_REFERENCE in conditions
    assert AWAY_FROM_HOME_REFERENCE in conditions
    evidence = " ".join(
        line.removeprefix("> ")
        for line in (PACKAGE / "evidence.md").read_text().splitlines()
    )
    evidence = " ".join(evidence.split())
    assert ACCOMMODATION_SUPPLEMENT_REFERENCE in evidence
    assert AWAY_FROM_HOME_REFERENCE in evidence
    for fact in _facts():
        assert fact.source.source_sha256 != PDF_SHA256
        assert fact.measure.concept_evidence_url == WORKBOOK_URL
        if fact.filters["benefit_expense_component"] == "Accommodation Assistance":
            assert "Accommodation Supplement" in fact.measure.concept_evidence_notes
            assert "Away from Home Allowance" in fact.measure.concept_evidence_notes
            assert fact.layout.groupby_value_label == "Accommodation Assistance"


def _guard_cases():
    """One affected spec for every distinct authored source guard/header."""
    cases = {}
    for index, spec in enumerate(_specs()):
        selector = spec.selector
        for address, message in (
            (selector.expected_row_header_address, "row header"),
            (selector.expected_column_header_address, "column header"),
            *((guard.address, guard.label) for guard in selector.guard_cells),
        ):
            if address is not None:
                cases.setdefault(address, (index, address, message))
    return tuple(cases.values())


GUARD_CASES = _guard_cases()


def _mutated_cells(address, *, raw_value="unexpected publisher layout"):
    assert sum(cell.address == address for cell in _cells()) == 1
    return [
        replace(cell, raw_value=raw_value) if cell.address == address else cell
        for cell in _cells()
    ]


def test_treasury_estimates_all_authored_guards_have_drift_regressions():
    assert {address for _, address, _ in GUARD_CASES} == {
        *(f"{column}1" for column in "BCEHKLMO"),
        *(f"{column}{row}" for row in range(2, 114) for column in "BCEHLMO"),
    }
    assert len(GUARD_CASES) == 792


@pytest.mark.parametrize(
    "spec_index,address,message",
    GUARD_CASES,
    ids=[address for _, address, _ in GUARD_CASES],
)
def test_treasury_estimates_each_guard_rejects_publisher_drift(
    spec_index, address, message
):
    with pytest.raises(ValueError, match=re.escape(message)):
        resolve_source_record(_mutated_cells(address), _specs()[spec_index])


def test_treasury_estimates_accommodation_assistance_source_rejects_as_relabelling():
    address = _compact_address(4783, "E")
    spec = next(
        spec
        for spec in _specs()
        if spec.selector.expected_row_header_address == address
    )
    with pytest.raises(ValueError, match="row header"):
        resolve_source_record(
            _mutated_cells(address, raw_value="Accommodation Supplement"), spec
        )


def test_treasury_estimates_accommodation_assistance_yaml_rejects_as_relabelling(
    tmp_path,
):
    payload = deepcopy(_payload())
    row = next(
        row
        for record_set in payload["record_sets"]
        for row in record_set["rows"]
        if row["expected_row_header"] == "Accommodation Assistance"
    )
    row["label"] = row["expected_row_header"] = "Accommodation Supplement"
    path = tmp_path / "source_package.yaml"
    path.write_text(yaml.safe_dump(payload, sort_keys=False))
    mutated = load_source_package(path)
    with pytest.raises(ValueError, match="row header"):
        mutated.build_facts(2026, cells=_cells(), source_rows=_rows())


@pytest.mark.parametrize("physical_row", range(4713, 4825))
def test_treasury_estimates_numeric_amount_type_guard_rejects_text(physical_row):
    address = _compact_address(physical_row)
    changed = [
        replace(cell, cell_type="text", raw_value="suppressed")
        if cell.address == address
        else cell
        for cell in _cells()
    ]
    spec = next(spec for spec in _specs() if spec.selector.address == address)
    with pytest.raises(ValueError, match="expected cell type 'number'"):
        resolve_source_record(changed, spec)


def _mutated_source_rows(mode):
    first_selected = next(row for row in _rows() if row.row_number == 4713)
    if mode == "absent":
        return [
            replace(row, values={**row.values, "App ID": -1})
            if row.row_number == first_selected.row_number
            else row
            for row in _rows()
        ]
    assert mode == "ambiguous"
    return [*_rows(), replace(first_selected, row_number=6453)]


@pytest.mark.parametrize("mode,count", (("absent", 0), ("ambiguous", 2)))
def test_treasury_estimates_selected_row_criteria_require_one_publisher_row(
    mode, count
):
    with pytest.raises(
        ValueError,
        match=f"Selected row criteria must match exactly one source row; .*matched {count} rows",
    ):
        _package().build_source_cells(2026, source_rows=_mutated_source_rows(mode))
