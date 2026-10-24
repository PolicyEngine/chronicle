"""Independent workbook checks for every selected RSZ quarterly data cell."""

from collections import Counter, defaultdict
import csv
from decimal import Decimal, ROUND_HALF_UP
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import re

from openpyxl import load_workbook
import pytest
import yaml

ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "packages/rsz/employment_daily_wage_class_2023_2026"
DATA = ROOT / "db/data/rsz/employment_daily_wage_class_2023_2026"
OLD = "rsz-arbeidsmarktanalyse-per-kwartaal-2023-2024.xlsx"
NEW = "rsz-arbeidsmarktanalyse-per-kwartaal-2024-2026.xlsx"


@pytest.fixture(scope="module")
def extract():
    with (DATA / "employment_daily_wage_class.csv").open(
        newline="", encoding="utf-8"
    ) as f:
        return list(csv.DictReader(f))


@pytest.fixture(scope="module")
def workbooks():
    result = {}
    for name in (OLD, NEW):
        result[name] = {}
        for formulas in (False, True):
            book = load_workbook(DATA / name, read_only=True, data_only=not formulas)
            try:
                for sheet in ("T17",) if name == OLD else ("T17", "T7"):
                    width = 7 if sheet == "T17" else 10
                    result[name][sheet, formulas] = [
                        tuple(row) for row in book[sheet].iter_rows(max_col=width)
                    ]
            finally:
                book.close()
    return result


def rounded(value, places):
    return Decimal(str(value)).quantize(
        Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP
    )


def test_every_ingested_cell_labels_and_display_precision(extract, workbooks):
    for name, sheets in workbooks.items():
        for sheet in ("T17",) if name == OLD else ("T17", "T7"):
            header = sheets[sheet, False][0 if sheet == "T17" else 1]
            offset = 4 if sheet == "T17" else 7
            assert [c.value for c in header[offset : offset + 3]] == [
                "Arbeidsplaatsen\nPostes de travail",
                "Tewerkgestelde werknemers\nTravailleurs occupés",
                "VTE\n|ETP",
            ]
    seen = set()
    count = Counter()
    noisy_samples = []
    age_types = set()
    unknown_printed = set()
    for row in extract:
        locator = json.loads(row["source_locator"])
        name, sheet, number = locator["workbook"], locator["sheet"], locator["row"]
        source = workbooks[name][sheet, False][number - 1]
        expressions = workbooks[name][sheet, True][number - 1]
        q = source[0].value
        assert row["quarter"] == f"{q // 10}Q{q % 10}"
        assert (q // 10 == 2023) if name == OLD else (2024 <= q // 10 <= 2026)
        offset = 4 if sheet == "T17" else 7
        wage = source[1 if sheet == "T17" else 6].value
        assert row["rsz.daily_wage_class"] == (
            "Onbekend-Inconnu" if wage.lower() == "onbekend-inconnu" else wage
        )
        if wage.lower() == "onbekend-inconnu":
            unknown_printed.add(wage)
        if sheet == "T17":
            assert row["rsz.age_class"] == str(source[2].value)
            if str(source[2].value) == "-20":
                age_types.add(type(source[2].value))
            assert row["rsz.sex"] == source[3].value
        else:
            assert row["rsz.worker_status"] == source[1].value
            assert row["rsz.sex"] == source[2].value
            assert row["rsz.nace2025_activity"] == source[3].value
            assert locator["printed_dimensions"][f"E{number}"] == source[4].value
            assert locator["printed_dimensions"][f"F{number}"] == source[5].value
        for column, metric in enumerate(("jobs", "employed_workers", "fte"), offset):
            cell = source[column]
            assert expressions[column].data_type != "f"
            fmt = cell.number_format
            # Restated publisher display contract, independently checked per cell.
            assert (
                "#,##0.0" in fmt
                if metric == "fte"
                else "#,##0" in fmt and "#,##0.0" not in fmt
            )
            expected = rounded(cell.value, 1 if metric == "fte" else 0)
            assert Decimal(row[metric]) == expected
            raw = locator["cells"][metric]
            assert raw == {
                "cell": cell.coordinate,
                "cached_value": str(cell.value),
                "number_format": fmt,
            }
            key = name, sheet, cell.coordinate
            assert key not in seen
            seen.add(key)
            count[sheet] += 1
            if metric == "fte" and abs(
                Decimal(str(cell.value)) - rounded(cell.value, 2)
            ) > Decimal("1e-13"):
                noisy_samples.append(key)
    assert count == {"T17": 5616, "T7": 29013}
    assert len(noisy_samples) > 100
    assert age_types == {int, str}
    assert unknown_printed == {"Onbekend-Inconnu", "Onbekend-inconnu"}
    assert {r["rsz.daily_wage_class"] for r in extract} == {
        "<100",
        "100 - 149",
        "150 - 199",
        "200 - 299",
        ">= 300",
        "Onbekend-Inconnu",
    }
    assert {r["rsz.age_class"] for r in extract if r["rsz.age_class"]} == {
        "-20",
        "20-24",
        "25-29",
        "30-34",
        "35-39",
        "40-44",
        "45-49",
        "50-54",
        "55-59",
        "60-64",
        "65-69",
        "70+",
    }
    # A specifically noisy cache is normalized at the actual one-decimal format.
    old_g2 = workbooks[OLD]["T17", False][1][6]
    assert old_g2.value == 10283.759999999998
    assert rounded(old_g2.value, 1) == Decimal("10283.8")


def test_no_missing_detail_cells_and_subtotal_checks(extract, workbooks):
    selected = {
        (
            json.loads(r["source_locator"])["workbook"],
            json.loads(r["source_locator"])["sheet"],
            json.loads(r["source_locator"])["row"],
        )
        for r in extract
    }
    expected = set()
    totals = 0
    display_differences = 0
    for name, sheets in workbooks.items():
        for (sheet, formulas), source in sheets.items():
            if formulas:
                continue
            offset = 4 if sheet == "T17" else 7
            for number, row in enumerate(source, 1):
                quarter = row[0].value
                if isinstance(quarter, int):
                    if name == NEW or quarter // 10 == 2023:
                        expected.add((name, sheet, number))
                elif str(quarter).startswith("Totaal"):
                    for column in range(offset, offset + 3):
                        formula = sheets[sheet, True][number - 1][column].value
                        match = re.fullmatch(
                            r"=SUBTOTAL\(9,([A-Z]+)(\d+):\1(\d+)\)", formula
                        )
                        assert match, formula
                        first, last = int(match[2]), int(match[3])
                        detail = [
                            source[n - 1][column].value for n in range(first, last + 1)
                        ]
                        assert all(isinstance(x, (int, float)) for x in detail)
                        # Cached cells sum exactly at their two-decimal cache representation.
                        # Display rounding is per fact, so it need not preserve this identity.
                        assert sum(
                            (rounded(x, 2) for x in detail), Decimal(0)
                        ) == rounded(row[column].value, 2)
                        if column == offset + 2:
                            display_differences += sum(
                                (rounded(x, 1) for x in detail), Decimal(0)
                            ) != rounded(row[column].value, 1)
                        totals += 1
                    assert (name, sheet, number) not in selected
    assert selected == expected
    assert totals == (8 + 9 + 9) * 3
    assert display_differences > 0


def test_full_2024_overlap_and_reported_display_differences(workbooks):
    def cells(name):
        result = {}
        for row in workbooks[name]["T17", False][1:]:
            q = row[0].value
            if isinstance(q, int) and q // 10 == 2024:
                key = (q, row[1].value.lower(), str(row[2].value), row[3].value.lower())
                for metric, cell in enumerate(row[4:]):
                    result[key + (metric,)] = cell
        return result

    older, newer = cells(OLD), cells(NEW)
    assert older.keys() == newer.keys()
    assert len(older) == 1728  # 4 quarters x 144 category rows x 3 measures
    assert sum(older[k].value != newer[k].value for k in older) == 285
    assert all(rounded(older[k].value, 2) == rounded(newer[k].value, 2) for k in older)
    differences = {
        (
            older[k].coordinate,
            newer[k].coordinate,
            str(rounded(older[k].value, 1)),
            str(rounded(newer[k].value, 1)),
        )
        for k in older
        if rounded(older[k].value, 1) != rounded(newer[k].value, 1)
    }
    report = json.loads((PACKAGE / "workbook_audit.json").read_text())
    recorded = {
        (d["older_cell"], d["newer_cell"], d["older_display"], d["newer_display"])
        for d in report["overlap"]["one_decimal_differences"]
    }
    assert differences == recorded
    assert len(differences) == 15


def test_declarations_and_distinct_measure_coverage(extract):
    payload = yaml.safe_load((PACKAGE / "source_package.yaml").read_text())
    assert len((PACKAGE / "source_package.yaml").read_bytes()) < 15_000_000
    assert not re.search(
        r"(^|\s)(?:&id\d+|\*id\d+|<<:)", (PACKAGE / "source_package.yaml").read_text()
    )
    groups = defaultdict(list)
    for record in payload["record_sets"]:
        assert record["provenance_class"] == "administrative"
        assert record["assertion"] == "observation"
        assert record["geography_name"] == "Belgium"
        assert len(record["measures"]) == 1
        metric = record["measures"][0]["measure_id"]
        quarter = record["period"]
        year, q = int(quarter[:4]), int(quarter[-1])
        expected_end = f"{year}-{q * 3:02d}-" + {1: "31", 2: "30", 3: "30", 4: "31"}[q]
        coverage = record["period_coverage"]
        assert coverage["end_date"] == expected_end
        assert coverage["start_date"] == (
            f"{year}-{q * 3 - 2:02d}-01" if metric == "fte" else expected_end
        )
        assert record["entity"] == "person"
        assert (
            record["entity_role"]
            == {
                "jobs": "employee_job",
                "employed_workers": "employee",
                "fte": "full_time_equivalent_employee",
            }[metric]
        )
        assert record["groupby_dimension"] in {"rsz.age_class", "rsz.nace2025_activity"}
        for row in record["rows"]:
            assert (
                row["label"]
                == payload["dimension_value_labels"][record["groupby_dimension"]][
                    row["value_id"]
                ]
            )
        groups[record["source_record_id_prefix"]].append(record)
    for records in groups.values():
        assert len(records) == 3
        assert len({r["record_set_spec_id"] for r in records}) == 1
        assert all(r["rows"] == records[0]["rows"] for r in records)
    assert sum(len(r["rows"]) for r in payload["record_sets"]) == 34629


def test_generator_reproduces_checked_outputs():
    spec = importlib.util.spec_from_file_location(
        "rsz_daily_generator", PACKAGE / "build_package.py"
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    rows, labels, variants = module.extract()
    csv_buffer = io.StringIO(newline="")
    writer = csv.DictWriter(csv_buffer, fieldnames=module.HEADERS, lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    assert (
        csv_buffer.getvalue() == (DATA / "employment_daily_wage_class.csv").read_text()
    )
    assert (
        module.yaml_text(module.package_spec(rows, labels))
        == (PACKAGE / "source_package.yaml").read_text()
    )
    assert variants == json.loads((PACKAGE / "nace_label_variants.json").read_text())
    manifest = yaml.safe_load((DATA / "manifest.yaml").read_text())
    for entry in manifest["files"].values():
        path = DATA / entry["filename"]
        assert hashlib.sha256(path.read_bytes()).hexdigest() == entry["sha256"]
        assert entry["storage"]["r2"]["key"].endswith(entry["sha256"] + "/" + path.name)
