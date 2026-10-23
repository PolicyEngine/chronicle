"""Literal-cell fidelity for Treasury PREFU 2026 earnings assumptions."""

from __future__ import annotations

import hashlib
import re
from dataclasses import replace
from decimal import Decimal
from functools import lru_cache
from pathlib import Path

import openpyxl
import pytest
import yaml

from chronicle.artifacts import build_r2_key
from chronicle.bundle import _dimension_label_reports, build_bundle_coverage
from chronicle.consumer_contract import (
    consumer_fact_rows,
    validate_consumer_fact_contract,
)
from chronicle.core import validate_facts
from chronicle.source_package import (
    SOURCE_PACKAGE_ALIASES,
    load_source_package,
    validate_source_package,
)
from chronicle.sources import build_source_cell_key, validate_source_cells
from chronicle.sources.specs import resolve_source_record

ROOT = Path(__file__).resolve().parents[1]
DIRECTORY = Path("treasury/prefu26_economic_assumptions")
PACKAGE = ROOT / "packages" / DIRECTORY
DATA = ROOT / "db/data" / DIRECTORY
ALIAS = "treasury-prefu26-economic-assumptions"
WORKBOOK = "prefu26-charts-data.xlsx"
SHA256 = "e2486375dde515d0722f8c3022f481804bb014beb83c81b8bfc99bc9426544d9"
URL = (
    "https://www.treasury.govt.nz/sites/default/files/2026-09/prefu26-charts-data.xlsx"
)
PAGE = "https://www.treasury.govt.nz/publications/efu/pre-election-economic-and-fiscal-update-2026"
SHEET = "Table 4.1"

# This independent literal map is not derived from source-package declarations.
# Footnote 4 qualifies D6: earnings are actual, although the column says Forecast4.
PUBLISHED_CELLS = (
    ("C17", 2025, "Actual", "4.2", "observation"),
    ("D17", 2026, "Forecast4", "2.9", "observation"),
    ("E17", 2027, "Forecast", "2.7", "source_projection"),
    ("F17", 2028, "Forecast", "3", "source_projection"),
    ("G17", 2029, "Forecast", "3.2", "source_projection"),
    ("H17", 2030, "Forecast", "3.4", "source_projection"),
    ("I17", 2031, "Forecast", "3.4", "source_projection"),
)


@lru_cache
def _package():
    return load_source_package(PACKAGE)


@lru_cache
def _cells():
    return _package().build_source_cells(2026)


@lru_cache
def _facts():
    return _package().build_facts(2026, cells=_cells())


@lru_cache
def _specs():
    return _package().build_source_record_specs(2026)


def _sheet_cells():
    return {cell.address: cell for cell in _cells() if cell.sheet_name == SHEET}


def test_prefu_alias_and_package_validation():
    assert SOURCE_PACKAGE_ALIASES[ALIAS] == DIRECTORY
    assert _package().artifact.artifact_year == 2026
    assert _package().artifact.sheets == ("Index", SHEET)
    report = validate_source_package(PACKAGE, year=2026)
    assert report.valid, report.to_dict()
    assert not report.warnings
    assert report.counts == {
        "record_set_count": 7,
        "row_count": 7,
        "measure_count": 7,
        "source_record_count": 7,
        "source_region_count": 7,
    }


def test_prefu_unchanged_workbook_receipt_license_and_canonical_r2_key():
    manifest = yaml.safe_load((DATA / "manifest.yaml").read_text())
    entry = manifest["files"][2026]
    content = (DATA / WORKBOOK).read_bytes()
    assert manifest["source_id"] == "treasury"
    assert manifest["package_id"] == ALIAS
    assert manifest["published_at"] == "2026-09-29"
    assert manifest["source_page"] == manifest["license_url"] == PAGE
    assert manifest["license"] == "CC BY 4.0"
    assert "Crown" in manifest["attribution"]
    assert entry["filename"] == WORKBOOK
    assert entry["source_url"] == URL
    assert entry["sha256"] == hashlib.sha256(content).hexdigest() == SHA256
    assert entry["size_bytes"] == len(content) == 664_670
    assert entry["fetched_at"] == "2026-10-22T07:57:29Z"
    key = build_r2_key(
        source_id="treasury",
        package_id=ALIAS,
        year=2026,
        sha256=SHA256,
        filename=WORKBOOK,
    )
    assert entry["storage"]["r2"] == {
        "provider": "r2",
        "bucket": "ledger-raw",
        "key": key,
        "uri": f"r2://ledger-raw/{key}",
    }


def test_prefu_complete_declared_sheet_ranges_and_publisher_footnotes():
    book = openpyxl.load_workbook(DATA / WORKBOOK, data_only=True, read_only=True)
    try:
        assert {cell.sheet_name for cell in _cells()} == {"Index", SHEET}
        for name in ("Index", SHEET):
            parsed = {
                cell.address: cell for cell in _cells() if cell.sheet_name == name
            }
            sheet = book[name]
            assert len(parsed) == sheet.max_row * sheet.max_column
            for row in sheet.iter_rows():
                for cell in row:
                    # ReadOnly EmptyCell has no coordinate; its grid address
                    # is nevertheless counted by the complete-range check.
                    if cell.value is not None:
                        assert parsed[cell.coordinate].raw_value == cell.value
    finally:
        book.close()
    assert validate_source_cells(_cells()).valid
    cells = _sheet_cells()
    assert cells["B17"].raw_value == "Average weekly earnings3 (ann % chg)"
    assert cells["C21"].raw_value == 3
    assert cells["D21"].raw_value == "Ordinary time."
    assert cells["C22"].raw_value == 4
    assert cells["D22"].raw_value == (
        "Measures are actuals except for real GDP and nominal GDP."
    )
    index = {cell.address: cell for cell in _cells() if cell.sheet_name == "Index"}
    assert index["A2"].raw_value == "Published 29 September 2026"
    assert "Creative Commons Attribution 4.0" in index["A13"].raw_value
    assert "Attribution to the Crown" in index["A14"].raw_value


@pytest.mark.parametrize("address,year,status,value,assertion", PUBLISHED_CELLS)
def test_prefu_every_published_earnings_cell_exactly_once(
    address, year, status, value, assertion
):
    cell = _sheet_cells()[address]
    assert Decimal(str(cell.raw_value)) == Decimal(value)
    facts = [fact for fact in _facts() if fact.period.value == year - 1]
    assert len(facts) == 1
    fact = facts[0]
    assert Decimal(str(fact.value)) == Decimal(value)
    assert fact.assertion == assertion
    assert fact.provenance_class == "model_output"
    assert fact.survey_instrument is None
    assert fact.period.type == "fiscal_year"
    assert fact.period_coverage.basis == "fiscal"
    assert fact.period_coverage.source_period_label == f"Year ending 30 June {year}"
    assert fact.period_coverage.start_date == f"{year - 1}-07-01"
    assert fact.period_coverage.end_date == f"{year}-06-30"
    assert "Ordinary time." in fact.period_coverage.notes
    assert "Measures are actuals except for real GDP and nominal GDP." in (
        fact.period_coverage.notes
    )
    assert fact.filters["earnings_component"] == "ordinary_time"
    assert fact.filters["publication_status"] == status
    assert {
        (constraint.variable, constraint.operator, constraint.value)
        for constraint in fact.constraints
    } == {
        ("earnings_component", "==", "ordinary_time"),
        ("publication_status", "==", status),
    }
    assert fact.measure.unit == "percent"
    assert fact.measure.source_concept == (
        "treasury.prefu26.average_weekly_earnings_annual_percentage_change"
    )
    assert build_source_cell_key(cell) in fact.source_cell_keys
    selected = [spec for spec in _specs() if spec.period == year - 1]
    assert len(selected) == 1
    assert selected[0].selector.address == address
    assert selected[0].selector.expected_row_header_address == "B17"
    assert selected[0].selector.expected_column_header_address == f"{address[0]}5"


def test_prefu_fact_boundary_labels_and_independent_coverage():
    assert len(_facts()) == 7
    assert validate_facts(_facts()).valid
    rows = consumer_fact_rows(_facts())
    assert validate_consumer_fact_contract(_facts()).valid
    assert _dimension_label_reports(ALIAS, rows) == ([], [])
    for fact in _facts():
        assert fact.source.source_name == "treasury"
        assert fact.geography.id == "NZ"
        assert fact.geography.name == "New Zealand"
        assert fact.entity.name == "government"
        assert fact.layout.groupby_dimension_label == "Economic assumption"
        assert fact.layout.groupby_value_label == (
            "Average weekly earnings (annual percentage change, ordinary time)"
        )
        assert "ordinary-time" in fact.layout.measure_label
        assert set(fact.dimension_labels) == {
            *fact.filters,
            fact.layout.groupby_dimension,
        }
        for dimension, value in fact.filters.items():
            assert fact.dimension_labels[dimension]
            assert fact.dimension_value_labels[dimension][str(value)]
    coverage = build_bundle_coverage(rows)
    assert coverage["fact_count"] == 7
    assert coverage["counts"]["by_source"] == {"treasury": 7}
    assert coverage["counts"]["by_geography"] == {"country:NZ": 7}
    assert coverage["counts"]["by_entity"] == {"government": 7}
    assert coverage["counts"]["by_period"] == {
        "fiscal_year:2024": 1,
        "fiscal_year:2025": 1,
        "fiscal_year:2026": 1,
        "fiscal_year:2027": 1,
        "fiscal_year:2028": 1,
        "fiscal_year:2029": 1,
        "fiscal_year:2030": 1,
    }
    assert coverage["duplicates"] == {
        "aggregate_fact_keys": [],
        "semantic_fact_keys": [],
    }


# Each authored guard is probed on a selected spec. Keep this map independent
# of the YAML so deleting a declaration makes its regression fail.
GUARD_CASES = (
    (0, "B17", "row header"),
    *((index, f"{column}5", "column header") for index, column in enumerate("CDEFGHI")),
    (0, "B1", "table title"),
    (0, "B2", "publisher source attribution"),
    (0, "B5", "publisher period basis"),
    *(
        (index, f"{column}6", "publisher status")
        for index, column in enumerate("CDEFGHI")
    ),
    (0, "C21", "ordinary-time footnote marker"),
    (0, "D21", "ordinary-time footnote"),
    (0, "C22", "actuals qualification marker"),
    (0, "D22", "actuals qualification"),
)


def test_prefu_every_authored_guard_has_an_independent_regression():
    declared = set()
    for spec in _specs():
        selector = spec.selector
        declared.add(selector.expected_row_header_address)
        declared.add(selector.expected_column_header_address)
        declared.update(guard.address for guard in selector.guard_cells)
    assert declared == {address for _, address, _ in GUARD_CASES}
    assert len(GUARD_CASES) == 22


@pytest.mark.parametrize(
    "spec_index,address,message",
    GUARD_CASES,
    ids=[address for _, address, _ in GUARD_CASES],
)
def test_prefu_each_guard_rejects_publisher_drift(spec_index, address, message):
    changed = [
        replace(cell, raw_value="unexpected publisher layout")
        if cell.sheet_name == SHEET and cell.address == address
        else cell
        for cell in _cells()
    ]
    with pytest.raises(ValueError, match=re.escape(message)):
        resolve_source_record(changed, _specs()[spec_index])


def test_prefu_row_shift_is_rejected_before_emitting_facts(tmp_path):
    payload = yaml.safe_load((PACKAGE / "source_package.yaml").read_text())
    payload["record_sets"][0]["rows"][0]["row_number"] = 16
    path = tmp_path / "source_package.yaml"
    path.write_text(yaml.safe_dump(payload, sort_keys=False))
    with pytest.raises(ValueError, match="row header"):
        load_source_package(path).build_facts(2026, cells=_cells())


@pytest.mark.parametrize("address", (cell[0] for cell in PUBLISHED_CELLS))
def test_prefu_each_selected_value_requires_a_numeric_publisher_cell(address):
    spec = next(spec for spec in _specs() if spec.selector.address == address)
    changed = [
        replace(cell, cell_type="text", raw_value="suppressed")
        if cell.sheet_name == SHEET and cell.address == address
        else cell
        for cell in _cells()
    ]
    with pytest.raises(ValueError, match="expected cell type 'number'"):
        resolve_source_record(changed, spec)


def test_prefu_build_year_does_not_relabel_published_reference_periods():
    for year in (2023, 2032):
        facts = _package().build_facts(year, cells=_cells())
        assert facts == _facts()
