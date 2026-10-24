"""Independent source-cell checks for RSZ annual wages, October 2026 release."""

from __future__ import annotations

from collections import Counter
import csv
from decimal import Decimal, ROUND_HALF_UP
import hashlib
from pathlib import Path
import runpy
import shutil

from openpyxl import load_workbook
import pytest
import yaml

from chronicle.consumer_contract import (
    consumer_fact_rows,
    validate_consumer_fact_contract,
)
from chronicle.core import validate_facts
from chronicle.source_package import load_source_package, validate_source_package
from chronicle.sources.cells import build_source_cell_key

ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "packages/rsz/annual_wages_key_figures_2020_2025"
DATA = ROOT / "db/data/rsz/annual_wages_key_figures_2020_2025"
RAW = DATA / "rsz-ruwedata-periods-wages.xlsx"
CSV = DATA / "annual_wages_key_figures_2020_2025.csv"
SHEET = "Kerncijfers-Chiffres Clés"


@pytest.fixture(scope="module")
def captured_rows():
    with CSV.open(encoding="utf-8", newline="") as stream:
        return list(csv.DictReader(stream))


@pytest.fixture(scope="module")
def annual():
    package = load_source_package(PACKAGE)
    source_rows = package.build_source_rows(2025)
    cells = package.build_source_cells(2025, source_rows=source_rows)
    facts = package.build_facts(2025, source_rows=source_rows, cells=cells)
    return package, cells, facts


def test_declarations_guards_and_calendar_coverage():
    report = validate_source_package(PACKAGE, year=2025)
    assert report.valid, report.to_dict()
    assert not report.warnings
    assert report.counts["record_set_count"] == 36
    assert report.counts["source_record_count"] == 252
    text = (PACKAGE / "source_package.yaml").read_text()
    payload = yaml.safe_load(text)
    assert not any(
        isinstance(token, (yaml.AnchorToken, yaml.AliasToken))
        for token in yaml.scan(text)
    )
    assert "<<:" not in text
    assert text.count("provenance_class: administrative") == 36
    assert payload["dimension_labels"] == {
        "rsz.annual_worker_type": "werknemerstype",
        "rsz.annual_sex": "Geslacht / Sexe",
        "rsz.annual_worker_type_sex": "werknemerstype / Geslacht / Sexe",
    }
    for record_set in payload["record_sets"]:
        year = record_set["period"]
        assert year in range(2020, 2026)
        assert record_set["provenance_class"] == "administrative"
        assert record_set["assertion"] == "observation"
        assert record_set["period_type"] == "calendar_year"
        assert record_set["period_coverage"]["start_date"] == f"{year}-01-01"
        assert record_set["period_coverage"]["end_date"] == f"{year}-12-31"
        assert record_set["period_coverage"]["basis"] == "calendar"
        assert record_set["period_coverage"]["source_period_label"] == str(year)
        notes = record_set["period_coverage"]["notes"]
        assert "pre-publication" in notes and "22/10/2026" in notes
        assert (
            "inbound cross-border" in notes
            and "posted" in notes
            and "self-employed" in notes
        )
        assert record_set["geography_id"] == "BE"
        assert record_set["geography_name"] == "Belgium"
        assert record_set["record_set_spec_id"] == "rsz.annual_wages.key_figures.v1"
        assert record_set["source_record_id_prefix"] == f"rsz.annual_wages.year{year}"
        assert len(record_set["measures"]) == 1
        for row in record_set["rows"]:
            guards = {g["label"]: g["expected_value"] for g in row["guard_cells"]}
            assert guards["workbook"] == RAW.name
            assert guards["sheet"] == SHEET
            assert guards["number_format"] == "#,##0.0"
            assert guards["cell"][0] in "DEFGHI"
            assert "value_scale" not in row
        assert "value_scale" not in record_set["measures"][0]


def test_capture_hashes_licence_receipts_and_r2_declarations():
    manifest = yaml.safe_load((DATA / "manifest.yaml").read_text())
    assert manifest["license"] == "CC BY 4.0"
    assert "RSZ/ONSS, Directie Statistiek" in manifest["attribution"]
    assert (
        manifest["files"]["source_capture"]["sha256"]
        == "7eea6aa74dd8c3c287c54bf591fef0e8e5d5507404988f94a05bf259fd6883af"
    )
    assert (DATA / "CAPTURED_AT.txt").read_text().strip() == "2026-10-24T01:05:27Z"
    checksums = dict(
        line.split()[::-1]
        for line in (DATA / "SHA256SUMS.txt").read_text().splitlines()
    )
    for key, entry in manifest["files"].items():
        path = DATA / entry["filename"]
        sha = hashlib.sha256(path.read_bytes()).hexdigest()
        assert sha == entry["sha256"]
        assert path.stat().st_size == entry["size_bytes"]
        assert entry["retrieved_at"] == "2026-10-24T01:05:27Z"
        expected_key = f"raw/belgium/rsz-annual-wages-key-figures-2020-2025/{key}/{sha}/{path.name}"
        assert entry["storage"]["r2"] == {
            "provider": "r2",
            "bucket": "ledger-raw",
            "key": expected_key,
            "uri": f"r2://ledger-raw/{expected_key}",
        }
        if path.name in checksums:
            assert sha == checksums[path.name]
    licence = (DATA / "rsz-open-data-data-governance-act.html").read_text()
    assert "CC BY 4.0" in licence
    assert "Wanneer u deze gegevens gebruikt" in licence


def test_every_selected_cell_matches_workbook_at_its_displayed_precision(captured_rows):
    """Expected layout and labels are restated here, independently of the generator."""
    workbook = load_workbook(RAW, data_only=True, read_only=True)
    try:
        sheet = workbook[SHEET]
        source = list(sheet.iter_rows())
        assert source[1][2].value == "Bezoldiging - rémunération (x 1.000 Eur)"
        expected_metrics = {
            "D": (
                "ordinary_remuneration",
                " gewone bezoldiging \nrémunération ordinaire",
            ),
            "E": ("severance", " verbrekingsvergoeding\nindemnité de rupture"),
            "F": ("premiums", " premies \nprimes"),
            "G": (
                "simple_departure_holiday_pay",
                " enkel vertrekvakantiegeld \npécule simple de sortie",
            ),
            "H": ("full_time_paid_days", " Dagen Voltijds \n Journées Temps Plein"),
            "I": (
                "part_time_paid_hours",
                " Uren  deeltijds \n Heures temps partiel\n(x 1.000)",
            ),
        }
        worker_labels = [
            "Arbeider - Ouvrier",
            "Bediende - Employé",
            "Ambtenaar - Fonctionnaire",
        ]
        sexes = ["Man - Homme", "Vrouw - Femme"]
        expected_coordinates = {
            f"{column}{number}" for number in range(4, 46) for column in "DEFGHI"
        }
        assert len(captured_rows) == 252
        assert Counter(row["cell"] for row in captured_rows) == Counter(
            expected_coordinates
        )
        for row in captured_rows:
            number = int(row["source_row"])
            column = row["source_column"]
            cell = source[number - 1][ord(column) - ord("A")]
            metric, heading = expected_metrics[column]
            year_index, within = divmod(number - 4, 7)
            year = 2020 + year_index
            assert row["workbook"] == RAW.name and row["sheet"] == SHEET
            assert row["cell"] == cell.coordinate
            assert row["year"] == str(year)
            assert row["measure_id"] == metric
            assert (
                row["measure_printed"]
                == source[2][ord(column) - ord("A")].value
                == heading
            )
            assert row["number_format"] == cell.number_format == "#,##0.0"
            assert row["raw_numeric_value"] == str(cell.value)
            expected = Decimal(str(cell.value)).quantize(
                Decimal("0.1"), rounding=ROUND_HALF_UP
            )
            assert row["value"] == format(expected, "f")
            assert Decimal(row["value"]).as_tuple().exponent == -1
            assert row["year_printed"] == (
                ""
                if source[number - 1][0].value is None
                else str(source[number - 1][0].value)
            )
            assert row["worker_type_printed"] == (source[number - 1][1].value or "")
            assert row["sex_printed"] == (source[number - 1][2].value or "")
            if within < 6:
                assert row["rsz.annual_worker_type"] == worker_labels[within // 2]
                assert row["rsz.annual_sex"] == sexes[within % 2]
            else:
                assert source[number - 1][0].value == f"Totaal {year}"
                assert row["rsz.annual_worker_type"] == row["rsz.annual_sex"] == ""
                assert row["rsz.annual_worker_type_sex"] == "total"
    finally:
        workbook.close()


def test_all_annual_total_cells_are_typed_values_and_remain_facts(captured_rows):
    workbook = load_workbook(RAW, data_only=False, read_only=True)
    try:
        rows = list(workbook[SHEET].iter_rows())
        expected = set()
        for year, row_number in zip(range(2020, 2026), (10, 17, 24, 31, 38, 45)):
            assert rows[row_number - 1][0].value == f"Totaal {year}"
            for cell in rows[row_number - 1][3:9]:
                assert cell.data_type == "n", (cell.coordinate, cell.value)
                assert isinstance(cell.value, (int, float))
                expected.add(cell.coordinate)
        actual = {
            row["cell"]
            for row in captured_rows
            if row["rsz.annual_worker_type_sex"] == "total"
        }
        assert actual == expected and len(actual) == 36
        assert all(cell.data_type != "f" for row in rows[3:] for cell in row[3:9])
    finally:
        workbook.close()


def test_binary_noise_is_removed_using_one_displayed_decimal(captured_rows):
    values = {row["cell"]: row for row in captured_rows}
    for address, raw_value, expected in (
        ("D4", "21462847.451270003", "21462847.5"),
        ("D25", "27202447.935229994", "27202447.9"),
        ("I39", "276000.22620999994", "276000.2"),
        ("E42", "816847.5043700002", "816847.5"),
    ):
        assert values[address]["raw_numeric_value"] == raw_value
        assert values[address]["number_format"] == "#,##0.0"
        assert values[address]["value"] == expected


def test_every_fact_keeps_value_units_coverage_and_publisher_locator(
    annual, captured_rows
):
    _, cells, facts = annual
    assert len(facts) == 252
    assert validate_facts(facts).valid
    exported = consumer_fact_rows(facts)
    assert validate_consumer_fact_contract(exported).valid
    assert Counter(fact.period.value for fact in facts) == {
        year: 42 for year in range(2020, 2026)
    }
    expected_units = {
        "ordinary_remuneration": "EUR_thousands",
        "severance": "EUR_thousands",
        "premiums": "EUR_thousands",
        "simple_departure_holiday_pay": "EUR_thousands",
        "full_time_paid_days": "day",
        "part_time_paid_hours": "hour_thousands",
    }
    by_id = {fact.source_record_id: fact for fact in facts}
    cells_by_key = {build_source_cell_key(cell): cell for cell in cells}
    assert len(by_id) == len(facts)
    for row in captured_rows:
        identifier = f"rsz.annual_wages.year{row['year']}.{row['rsz.annual_worker_type_sex']}.{row['measure_id']}"
        fact = by_id[identifier]
        assert Decimal(str(fact.value)) == Decimal(row["value"])
        assert fact.measure.unit == expected_units[row["measure_id"]]
        assert (
            fact.provenance_class == "administrative"
            and fact.assertion == "observation"
        )
        assert fact.period_coverage.start_date == f"{row['year']}-01-01"
        assert fact.period_coverage.end_date == f"{row['year']}-12-31"
        assert fact.geography.name == "Belgium"
        assert fact.layout.table_record_kind == (
            "total" if row["rsz.annual_worker_type_sex"] == "total" else "detail"
        )
        lineage_values = {
            str(cells_by_key[key].raw_value) for key in fact.source_cell_keys
        }
        assert RAW.name in lineage_values and SHEET in lineage_values
        assert row["cell"] in lineage_values and "#,##0.0" in lineage_values
        for dimension in ("rsz.annual_worker_type", "rsz.annual_sex"):
            if row[dimension]:
                assert fact.filters[dimension] == row[dimension]
            else:
                assert dimension not in fact.filters


def test_generator_reproduces_checked_bytes(tmp_path):
    generator = runpy.run_path(str(PACKAGE / "build_package.py"))
    data_dir, package_dir = tmp_path / "data", tmp_path / "package"
    data_dir.mkdir()
    manifest = yaml.safe_load((DATA / "manifest.yaml").read_text())
    for key, entry in manifest["files"].items():
        if key != 2025:
            shutil.copyfile(DATA / entry["filename"], data_dir / entry["filename"])
    generator["generate"](data_dir=data_dir, package_dir=package_dir)
    assert (data_dir / CSV.name).read_bytes() == CSV.read_bytes()
    assert (data_dir / "manifest.yaml").read_bytes() == (
        DATA / "manifest.yaml"
    ).read_bytes()
    assert (package_dir / "source_package.yaml").read_bytes() == (
        PACKAGE / "source_package.yaml"
    ).read_bytes()
