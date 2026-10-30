"""Source packages for PolicyEngine/chronicle#395.

Northern Ireland's domestic rates (LPS rates paid; ONS's regional and district rows),
Wales's bus journeys (DfT BUS01), the Childcare Offer for Wales take-up, the national
accounts' split of vehicle excise duty, the OBR's VED line in both EFO vintages, and the
CPI weights and indices for gas and its two classes.
"""

from __future__ import annotations

from collections import Counter
import csv
import hashlib
import io
from pathlib import Path
import runpy

from openpyxl import load_workbook
import pandas as pd
import pytest
import yaml

from chronicle.bundle import UK_BUNDLE_SOURCES
from chronicle.consumer_contract import (
    build_semantic_fact_key,
    validate_consumer_fact_contract,
)
from chronicle.core import validate_facts
from chronicle.source_package import (
    SOURCE_PACKAGE_ALIASES,
    load_source_package,
    validate_source_package,
)
from chronicle.sources import cells, rows

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "db/data"
# package id -> (artifact year, record sets, facts)
PACKAGES = {
    "lps-domestic-rates-paid-2026": (2026, 6, 144),
    "ons-crpsf-ni-domestic-rates-fye2025": (2025, 52, 52),
    "dft-bus01-wales-journeys-2025": (2025, 57, 57),
    "welshgov-childcare-offer-report-2023-24": (2025, 3, 3),
    "ons-pusf-vehicle-duty-households-october-2026": (2026, 8, 8),
    "ons-pusf-vehicle-duty-businesses-october-2026": (2026, 8, 8),
    "ons-mm23-cpi-weight-gas-october-2026": (2026, 3, 3),
    "ons-mm23-cpi-weight-natural-gas-october-2026": (2026, 3, 3),
    "ons-mm23-cpi-weight-liquefied-hydrocarbons-october-2026": (2026, 3, 3),
    "ons-mm23-cpi-index-natural-gas-october-2026": (2026, 7, 7),
    "ons-mm23-cpi-index-liquefied-hydrocarbons-october-2026": (2026, 7, 7),
}


def _facts(package_id: str) -> dict[str, object]:
    year = PACKAGES[package_id][0] if package_id in PACKAGES else 2026
    package = load_source_package(package_id)
    return {fact.source_record_id: fact for fact in package.build_facts(year)}


def _series(cdid: str, directory: str) -> dict[str, float]:
    """The series' period rows ('2025', '2025 Q3', '2025 JAN'), past its metadata rows."""
    text = (DATA / directory / f"{cdid}.csv").read_text(encoding="utf-8")
    return {
        label: float(value)
        for label, value in csv.reader(io.StringIO(text))
        if label[:1].isdigit() and value
    }


@pytest.mark.parametrize("package_id", sorted(PACKAGES))
def test_packages_validate_register_and_carry_no_anchors(package_id):
    year, record_sets, count = PACKAGES[package_id]
    package_dir = ROOT / "packages" / SOURCE_PACKAGE_ALIASES[package_id]
    report = validate_source_package(package_dir, year=year)
    assert report.valid, report.to_dict()
    assert not report.warnings
    assert report.counts["record_set_count"] == record_sets
    assert package_id in UK_BUNDLE_SOURCES
    text = (package_dir / "source_package.yaml").read_text(encoding="utf-8")
    assert not any(
        isinstance(token, (yaml.AnchorToken, yaml.AliasToken))
        for token in yaml.scan(text)
    )
    built = list(_facts(package_id).values())
    assert len(built) == count
    assert validate_facts(built).valid
    assert validate_consumer_fact_contract(built).valid
    assert max(Counter(build_semantic_fact_key(fact) for fact in built).values()) == 1


def test_generator_reproduces_the_committed_packages():
    generator = runpy.run_path(str(ROOT / "scripts/build_uk_395_packages.py"))
    for directory, build in generator["BUILDERS"].items():
        path = ROOT / "packages" / directory / "source_package.yaml"
        assert build() == path.read_text(encoding="utf-8"), directory


@pytest.mark.parametrize(
    "package_id", sorted(set(PACKAGES) - {"dft-bus01-wales-journeys-2025"})
)
def test_manifests_pin_bytes_licence_and_r2(package_id):
    data_dir = DATA / SOURCE_PACKAGE_ALIASES[package_id]
    manifest = yaml.safe_load((data_dir / "manifest.yaml").read_text(encoding="utf-8"))
    assert manifest["package_id"] == package_id
    assert manifest["license"] == "Open Government Licence v3.0"
    assert manifest["license_verified_at"] == "2026-10-30"
    (entry,) = manifest["files"].values()
    content = (data_dir / entry["filename"]).read_bytes()
    assert hashlib.sha256(content).hexdigest() == entry["sha256"]
    assert len(content) == entry["size_bytes"]
    assert entry["storage"]["r2"]["uri"].endswith(
        f"/{package_id}/{PACKAGES[package_id][0]}/{entry['sha256']}/{entry['filename']}"
    )


def test_the_delimited_parser_reads_pound_amounts():
    """LPS prints amounts as '£51,877,954'; '£' is stripped as '$' is, and only from
    numbers, so a band label keeps its text."""
    for scalar in (rows._delimited_scalar, cells._delimited_scalar):
        assert scalar("£51,877,954") == 51877954
        assert scalar("-£1,234") == -1234
        assert scalar("$1,234") == 1234
        assert scalar("64,764 ") == 64764
        assert scalar("£0.01 - £10.00") == "£0.01 - £10.00"
        assert scalar("£") == "£"


def test_lps_facts_are_the_csv_cells():
    path = (
        DATA
        / "lps/domestic_rates_paid_2026/domestic-rates-collected-by-district-council.csv"
    )
    table = list(csv.reader(io.StringIO(path.read_text(encoding="utf-8-sig"))))
    header, body = table[0], table[1:]
    facts = _facts("lps-domestic-rates-paid-2026")
    for opening in range(2020, 2026):
        label = f"{opening}-{(opening + 1) % 100:02d}"
        count = header.index(
            f"Number of Domestic occupancies with a rates payment made in {label} rating year"
        )
        paid = header.index(f"Domestic Rates paid as at 31st March {opening + 1}")
        for row in body:
            value_id = row[0].lower().replace(",", "").replace(" ", "_")
            prefix = f"lps.domestic_rates_paid_2026.fy{opening}.{value_id}"
            assert facts[f"{prefix}.occupancies_with_payment"].value == int(
                row[count].strip().replace(",", "")
            )
            assert facts[f"{prefix}.domestic_rates_paid"].value == int(
                row[paid].replace("£", "").replace(",", "")
            )
        # The councils' counts sum to Northern Ireland's; the amounts to within the
        # publisher's rounding (no more than GBP 2).
        councils = [
            fact
            for key, fact in facts.items()
            if key.startswith(f"lps.domestic_rates_paid_2026.fy{opening}.")
            and fact.geography.level == "local_authority"
        ]
        assert len(councils) == 22
        for measure in ("occupancies_with_payment", "domestic_rates_paid"):
            total = facts[
                f"lps.domestic_rates_paid_2026.fy{opening}.northern_ireland.{measure}"
            ].value
            parts = sum(
                f.value for f in councils if f.source_record_id.endswith(measure)
            )
            assert abs(parts - total) <= (
                0 if measure == "occupancies_with_payment" else 2
            )
    # The 2024-25 total PolicyEngine/policyengine-uk#2663 cites.
    assert (
        facts[
            "lps.domestic_rates_paid_2026.fy2024.northern_ireland.domestic_rates_paid"
        ].value
        == 889543160
    )


def test_crpsf_rows_are_the_northern_ireland_sheet():
    workbook = load_workbook(
        DATA
        / "ons/crpsf_ni_domestic_rates_fye2025/crpsffye2025totalpublicsectorcurrentreceipts.xlsx",
        data_only=True,
    )
    sheet, united_kingdom = workbook["Northern Ireland"], workbook["United Kingdom"]
    facts = _facts("ons-crpsf-ni-domestic-rates-fye2025")
    for row, line_id in (
        (51, "ni_regional_domestic_rates"),
        (52, "ni_district_domestic_rates"),
    ):
        for column in range(2, 28):
            opening = int(sheet.cell(3, column).value.split(" to ")[0])
            fact = facts[f"ons.crpsf_fye2025.{line_id}.fy{opening}.{line_id}.amount"]
            assert fact.value == sheet.cell(row, column).value * 1_000_000
            # Every Northern Ireland domestic rate is Northern Ireland's.
            assert (
                sheet.cell(row, column).value == united_kingdom.cell(row, column).value
            )
    assert (
        facts[
            "ons.crpsf_fye2025.ni_regional_domestic_rates.fy2024.ni_regional_domestic_rates.amount"
        ].value
        == 498_000_000
    )
    assert (
        facts[
            "ons.crpsf_fye2025.ni_district_domestic_rates.fy2024.ni_district_domestic_rates.amount"
        ].value
        == 441_000_000
    )


def test_bus01_wales_mirrors_the_england_package():
    workbook = DATA / "dft/bus01_passenger_journeys_2025/bus01.ods"
    bus01a = pd.read_excel(workbook, sheet_name="BUS01a", header=None, engine="odf")
    bus01c = pd.read_excel(workbook, sheet_name="BUS01c", header=None, engine="odf")
    assert bus01a.iloc[7, 17] == "Wales" and bus01c.iloc[7, 15] == "Wales"
    facts = _facts("dft-bus01-wales-journeys-2025")
    england = list(_facts_at("dft-bus01-passenger-journeys-2025", 2025).values())
    assert {fact.measure.concept for fact in facts.values()} == {
        fact.measure.concept for fact in england
    }
    for index in range(8, len(bus01a)):
        year = bus01a.iloc[index, 0]
        if isinstance(year, (int, float)) and 2005 <= year <= 2025:
            key = f"dft.bus01.wales.bus01a.passenger_journeys.fy{int(year)}"
            fact = facts[f"{key}.wales.passenger_journeys"]
            assert fact.value == pytest.approx(bus01a.iloc[index, 17] * 1_000_000)
    categories = {
        "Elderly/Disabled journeys": "elderly_disabled_concessionary",
        "Total concessionary journeys": "total_concessionary",
    }
    found = 0
    for index in range(8, len(bus01c)):
        year, kind = bus01c.iloc[index, 0], bus01c.iloc[index, 1]
        if kind in categories and 2008 <= year <= 2025:
            key = f"dft.bus01.wales.bus01c.{categories[kind]}.fy{int(year)}"
            fact = facts[f"{key}.wales.passenger_journeys"]
            assert fact.value == pytest.approx(bus01c.iloc[index, 15] * 1_000_000)
            assert dict(fact.filters) == {"journey_category": categories[kind]}
            found += 1
    assert found == 36
    assert all(fact.geography.id == "W92000004" for fact in facts.values())


def _facts_at(package_id: str, year: int) -> dict[str, object]:
    package = load_source_package(package_id)
    return {fact.source_record_id: fact for fact in package.build_facts(year)}


def test_childcare_offer_table_4():
    facts = _facts("welshgov-childcare-offer-report-2023-24")
    prefix = "welshgov.childcare_offer_report_2023_24.table_4"
    expected = {
        "children_in_offer_funded_childcare": (21283, "administrative", "count"),
        "estimated_eligible_population": (36600, "model_output", "count"),
        "estimated_take_up_rate": (58, "model_output", "percent"),
    }
    for value_id, (value, provenance, unit) in expected.items():
        fact = facts[f"{prefix}.{value_id}.ay2023.{value_id}.{value_id}"]
        assert fact.value == value
        assert fact.provenance_class == provenance
        assert fact.measure.unit == unit
        assert (fact.period.type, fact.period.value) == ("academic_year", 2023)
        assert (fact.period_coverage.start_date, fact.period_coverage.end_date) == (
            "2023-09-01",
            "2024-08-31",
        )


def test_vehicle_duty_split_is_the_series():
    for cdid, side in (("cddz", "households"), ("eked", "businesses")):
        directory = f"ons/pusf_vehicle_duty_{side}_october_2026"
        series = _series(cdid, directory)
        facts = _facts(f"ons-pusf-vehicle-duty-{side}-october-2026")
        assert len(facts) == 8
        for year in range(2018, 2026):
            fact = facts[f"ons.pusf_{cdid}.cy{year}.{cdid}.value"]
            assert fact.value == series[str(year)] * 1_000_000
            assert fact.measure.concept == f"ons.vehicle_duty_paid_by_{side}"


def test_gas_weights_add_up_and_indices_are_published_averages():
    weights = {
        cdid: _series(cdid, f"ons/mm23_cpi_weight_{name}_october_2026")
        for cdid, name in (
            ("cjxb", "gas"),
            ("l89x", "natural_gas"),
            ("l89y", "liquefied_hydrocarbons"),
        )
    }
    for year in ("2024", "2025", "2026"):
        assert weights["l89x"][year] + weights["l89y"][year] == pytest.approx(
            weights["cjxb"][year], abs=1e-4
        )
    facts = _facts("ons-mm23-cpi-weight-liquefied-hydrocarbons-october-2026")
    assert facts["ons.mm23_l89y.cy2026.l89y.value"].value == 0.3326
    assert facts["ons.mm23_l89y.cy2026.l89y.value"].measure.unit == "parts_per_thousand"
    for cdid, name in (("l7e9", "natural_gas"), ("l7ea", "liquefied_hydrocarbons")):
        series = _series(cdid, f"ons/mm23_cpi_index_{name}_october_2026")
        months = (
            "JAN",
            "FEB",
            "MAR",
            "APR",
            "MAY",
            "JUN",
            "JUL",
            "AUG",
            "SEP",
            "OCT",
            "NOV",
            "DEC",
        )
        facts = _facts(f"ons-mm23-cpi-index-{name.replace('_', '-')}-october-2026")
        for year in (2023, 2024, 2025):
            mean = sum(series[f"{year} {month}"] for month in months) / 12
            assert series[str(year)] == pytest.approx(mean, abs=0.051)
            assert (
                facts[f"ons.mm23_{cdid}.cy{year}.{cdid}.value"].value
                == series[str(year)]
            )
        for year in (2023, 2024, 2025, 2026):
            fact = facts[f"ons.mm23_{cdid}.m{year}_01.{cdid}.value"]
            assert fact.value == series[f"{year} JAN"]
            assert (fact.period.type, fact.period.value) == ("month", f"{year}-01")


def test_ved_is_one_concept_in_both_efo_vintages():
    march = _facts_at("obr-efo-receipts-march-2026", 2026)
    october = _facts_at("obr-efo-receipts-october-2026", 2026)
    march_ved = {
        key: fact
        for key, fact in march.items()
        if key.startswith("obr.efo_2026_03.receipts.vehicle_excise_duties.")
    }
    october_ved = {
        key: fact
        for key, fact in october.items()
        if key.startswith("obr.efo_2026_10.receipts.vehicle_excise_duties.")
    }
    assert len(march_ved) == 7 and len(october_ved) == 7
    for fact in [*march_ved.values(), *october_ved.values()]:
        assert fact.measure.concept == "obr.vehicle_excise_duties"
        assert dict(fact.filters) == {"accounting_basis": "cash"}
    outturn = march_ved[
        "obr.efo_2026_03.receipts.vehicle_excise_duties.fy2024.vehicle_excise_duties.amount"
    ]
    assert outturn.value == 8.205e9
    assert outturn.assertion == "observation"
    assert not any("cash_receipts.vehicle_excise_duties" in key for key in october)
