"""Independent publisher-cell checks for PensionStat's 2025 workbook slice."""

from __future__ import annotations

from collections import Counter
import csv
from dataclasses import replace
import hashlib
import importlib.util
import json
from pathlib import Path
import re
import sys
import unicodedata

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter
import pytest
import yaml

from chronicle.core import ALLOWED_PERIOD_BASES
from chronicle.source_package import load_source_package
from chronicle.sources.cells import SourceArtifactMetadata, SourceCell
from chronicle.sources.specs import resolve_source_record
from chronicle.yaml_io import safe_load


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "db/data/sfpd/legal_pension_caseload_2025"
PACKAGE_DIR = ROOT / "packages/sfpd/legal_pension_caseload_2025"
RAW = DATA_DIR / "data-2025-nl.xlsx"
PACKAGE_ID = "sfpd-legal-pension-caseload-2025"
SHA256 = "d358fff0c643c8edfd53b55514c311d7f98a03e3c7d5a37be1450e5ce84f94a2"
SOURCE_URL = (
    "https://www.pensionstat.be/file/6103c41969844d308979992d7f84d51e59927e1e/"
    "b7bc512baaae29d4344ac44519e892339ddf5d19/data-2025-nl.xlsx"
)
# Columns are the publisher's native value columns, independently specified here.
VALUE_COLUMNS = {
    "Gepensioneerden": {"I": "sum", "J": "mean"},
    "Gepensioneerden_gemeente": {"C": "sum"},
    "IGO": {"I": "sum", "J": "mean", "K": "median"},
    "Jaaruitgave": {"G": "sum"},
    "MinP_aantal": {"I": "sum"},
}


@pytest.fixture(scope="module")
def package():
    return load_source_package(PACKAGE_ID)


@pytest.fixture(scope="module")
def specs(package):
    return package.build_source_record_specs(2025)


@pytest.fixture(scope="module")
def publisher_rows():
    """Read publisher bytes independently, retaining only headers and 2025 rows."""
    workbook = openpyxl.load_workbook(RAW, read_only=True, data_only=True)
    found = {}
    try:
        for sheet_name in VALUE_COLUMNS:
            rows = {}
            for row_number, values in enumerate(
                workbook[sheet_name].iter_rows(values_only=True), start=1
            ):
                if row_number == 1 or values[0] == 2025:
                    rows[row_number] = values
            assert rows[1][0] == "Jaar"
            assert len(rows) > 1
            found[sheet_name] = rows
    finally:
        workbook.close()
    return found


@pytest.fixture(scope="module")
def artifact(package):
    manifest = yaml.safe_load((DATA_DIR / "manifest.yaml").read_text())
    entry = manifest["files"][2025]
    r2 = entry["storage"]["r2"]
    return SourceArtifactMetadata(
        source_name=package.artifact.source_name,
        source_table=package.artifact.source_table,
        source_file=entry["filename"],
        url=entry["source_url"],
        vintage=package.artifact.vintage,
        sha256=entry["sha256"],
        size_bytes=entry["size_bytes"],
        extracted_at=package.artifact.extracted_at,
        extraction_method=package.artifact.extraction_method,
        raw_r2_bucket=r2["bucket"],
        raw_r2_key=r2["key"],
        raw_r2_uri=r2["uri"],
    )


def _source_value(rows, address):
    match = re.fullmatch(r"([A-Z]+)([1-9][0-9]*)", address)
    assert match is not None
    column, row = match.groups()
    return rows[int(row)][column_index_from_string(column) - 1]


def _sample_cells(spec, rows, artifact):
    """Construct only the independently read cells this one selector references."""
    selector = spec.selector
    addresses = {
        selector.address,
        selector.expected_row_header_address,
        selector.expected_column_header_address,
        *(guard.address for guard in selector.guard_cells),
    }
    cells = []
    for address in sorted(address for address in addresses if address is not None):
        value = _source_value(rows, address)
        match = re.fullmatch(r"([A-Z]+)([0-9]+)", address)
        assert match is not None
        column, row = match.groups()
        cells.append(
            SourceCell(
                artifact=artifact,
                sheet_name=selector.sheet_name,
                row_number=int(row),
                column_number=column_index_from_string(column),
                address=address,
                cell_type="number" if isinstance(value, (int, float)) else "text",
                raw_value=value,
                display_value=None if value is None else str(value),
            )
        )
    return cells


def test_pensionstat_raw_artifact_identity_and_storage():
    entry = yaml.safe_load((DATA_DIR / "manifest.yaml").read_text())["files"][2025]
    content = RAW.read_bytes()
    assert entry["filename"] == RAW.name
    assert len(content) == entry["size_bytes"] == 5_939_528
    assert hashlib.sha256(content).hexdigest() == entry["sha256"] == SHA256
    assert entry["source_url"] == SOURCE_URL
    r2 = entry["storage"]["r2"]
    assert r2["bucket"] == "ledger-raw"
    assert r2["key"] == f"raw/belgium/{PACKAGE_ID}/2025/{SHA256}/{RAW.name}"
    assert r2["uri"] == f"r2://ledger-raw/{r2['key']}"


def test_pensionstat_generator_reproduces_explicit_checked_package():
    path = PACKAGE_DIR / "build_package.py"
    module_spec = importlib.util.spec_from_file_location("pensionstat_generator", path)
    assert module_spec is not None and module_spec.loader is not None
    module = importlib.util.module_from_spec(module_spec)
    sys.modules[module_spec.name] = module
    try:
        module_spec.loader.exec_module(module)
        generated = module.generate_package()
    finally:
        sys.modules.pop(module_spec.name, None)
    text = (PACKAGE_DIR / "source_package.yaml").read_text()
    checked = safe_load(text)
    assert generated == checked
    assert all(
        record_set["provenance_class"] == "administrative"
        for record_set in checked["record_sets"]
    )
    assert text.count("provenance_class:") == len(checked["record_sets"])
    assert not any(
        isinstance(token, (yaml.tokens.AnchorToken, yaml.tokens.AliasToken))
        for token in yaml.scan(text, Loader=yaml.CSafeLoader)
    )
    assert not re.search(r"(?m)^\s*<<:", text)


def test_pensionstat_covers_each_published_2025_value_once(specs, publisher_rows):
    expected = set()
    for sheet_name, columns in VALUE_COLUMNS.items():
        for row_number, values in publisher_rows[sheet_name].items():
            if row_number == 1:
                continue
            for column in columns:
                value = values[column_index_from_string(column) - 1]
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    expected.add((sheet_name, f"{column}{row_number}"))
    selected = Counter(
        (spec.selector.sheet_name, spec.selector.address) for spec in specs
    )
    assert set(selected) == expected
    assert set(selected.values()) == {1}
    assert len({spec.source_record_id for spec in specs}) == len(specs)
    for spec in specs:
        assert spec.provenance_class == "administrative"
        assert spec.period_coverage is not None
        assert spec.period_coverage.basis in ALLOWED_PERIOD_BASES
        assert spec.value_scale == 1
        assert spec.divisor_selector is None
        assert spec.round_to is None
        assert spec.selector.end_address is None


@pytest.mark.parametrize("sheet_name", VALUE_COLUMNS)
def test_pensionstat_sample_cells_resolve_exact_publisher_values(
    sheet_name, specs, publisher_rows, artifact
):
    rows = publisher_rows[sheet_name]
    for column, aggregation in VALUE_COLUMNS[sheet_name].items():
        matching = [
            spec
            for spec in specs
            if spec.selector.sheet_name == sheet_name
            and spec.selector.address.startswith(column)
        ]
        assert matching
        for position in sorted({0, len(matching) // 2, len(matching) - 1}):
            spec = matching[position]
            selector = spec.selector
            cells = _sample_cells(spec, rows, artifact)
            record = resolve_source_record(cells, spec)
            assert record.value == _source_value(rows, selector.address)
            assert selector.address in record.source_cell_addresses
            assert len(record.source_cell_keys) == len(record.source_cell_addresses)
            assert spec.aggregation == aggregation
            if sheet_name == "Jaaruitgave":
                assert (spec.period_type, str(spec.period)) == ("calendar_year", "2025")
            else:
                assert (spec.period_type, spec.period) == ("month", "2025-01")
            row_number = int(re.search(r"[0-9]+$", selector.address).group())
            assert any(
                guard.address == f"A{row_number}" and guard.expected_value == 2025
                for guard in selector.guard_cells
            )
            for dimension_column, header in enumerate(rows[1], start=1):
                if (
                    dimension_column == 1
                    or get_column_letter(dimension_column) in VALUE_COLUMNS[sheet_name]
                ):
                    continue
                dimension = "pensionstat." + re.sub(
                    r"[^a-z0-9]+", "_", header.lower()
                ).strip("_")
                value = rows[row_number][dimension_column - 1]
                if value == "":
                    # Blank expenditure tariffs stay guarded without inventing
                    # a dimension category absent from the publisher's table.
                    assert dimension not in spec.filters
                    continue
                assert spec.filters[dimension] == value
                assert any(
                    constraint.variable == dimension
                    and constraint.operator == "=="
                    and constraint.value == value
                    for constraint in spec.constraints
                )


def test_pensionstat_dimension_guard_rejects_shifted_source(
    specs, publisher_rows, artifact
):
    spec = next(spec for spec in specs if spec.selector.sheet_name == "Gepensioneerden")
    cells = _sample_cells(spec, publisher_rows["Gepensioneerden"], artifact)
    guard = next(
        guard for guard in spec.selector.guard_cells if guard.address.startswith("B")
    )
    changed = [
        replace(cell, raw_value="shifted publisher category")
        if cell.address == guard.address
        else cell
        for cell in cells
    ]
    with pytest.raises(ValueError):
        resolve_source_record(changed, spec)


def test_pensionstat_pinned_workbook_keeps_publisher_periods(package, specs):
    assert package.artifact.artifact_year == 2025
    assert package.artifact.parser == "xlsx_read_only_used_range"
    assert set(package.artifact.sheets) == {"Info", *VALUE_COLUMNS}
    assert package.build_source_record_specs(2023) == specs


FR_RAW = DATA_DIR / "data-2025-fr.xlsx"
FR_SHA256 = "af638bd6dc344e83b514d91c7c2cb0b5b4cd00a8da10f8659ac452fbe8b79c33"
STATBEL_COMMUNES = (
    ROOT
    / "db/data/statbel/population_structure_commune_2025"
    / "statbel_population_structure_commune_2025.csv"
)


def _orthography(name):
    text = name.casefold().replace("œ", "oe").replace("’", "'")
    return "".join(
        character
        for character in unicodedata.normalize("NFKD", text)
        if not unicodedata.combining(character)
    )


def _commune_sheet(path, sheet):
    """Every row of a commune sheet, read independently of the generator."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return list(workbook[sheet].iter_rows(values_only=True))
    finally:
        workbook.close()


def test_pensionstat_french_label_evidence_identity_and_storage():
    entry = yaml.safe_load((DATA_DIR / "manifest.yaml").read_text())["files"][
        "label_evidence_fr"
    ]
    content = FR_RAW.read_bytes()
    assert entry["filename"] == FR_RAW.name
    assert len(content) == entry["size_bytes"] == 5_621_010
    assert hashlib.sha256(content).hexdigest() == entry["sha256"] == FR_SHA256
    r2 = entry["storage"]["r2"]
    assert r2["bucket"] == "ledger-raw"
    assert r2["key"] == (
        f"raw/belgium/{PACKAGE_ID}/label_evidence_fr/{FR_SHA256}/{FR_RAW.name}"
    )
    assert r2["uri"] == f"r2://ledger-raw/{r2['key']}"


def test_pensionstat_communes_biject_onto_nis_2025_with_evidence(specs):
    statbel = {}
    with STATBEL_COMMUNES.open(newline="", encoding="utf-8") as stream:
        for row in csv.DictReader(stream):
            if row["geography_level"] == "commune":
                statbel[row["geography_id"]] = (
                    row["geography_name"],
                    row["source_municipality_nl"],
                    row["source_municipality_fr"],
                )
    assert len(statbel) == 565
    nl_rows = _commune_sheet(RAW, "Gepensioneerden_gemeente")
    fr_rows = _commune_sheet(FR_RAW, "Pensionnés_commune")
    assert nl_rows[0] == ("Jaar", "Gemeente", "Aantal")
    assert fr_rows[0] == ("Année", "Commune", "Nombre")

    def series(rows):
        found = {}
        for year, name, count in rows[1:]:
            found.setdefault(name, {})[year] = count
        return {name: tuple(sorted(counts.items())) for name, counts in found.items()}

    nl_series, fr_series = series(nl_rows), series(fr_rows)
    audit = json.loads((PACKAGE_DIR / "commune_mapping.json").read_text())
    by_cell = {entry["source_cell"]: entry for entry in audit}
    commune_specs = [
        spec for spec in specs if spec.selector.sheet_name == "Gepensioneerden_gemeente"
    ]
    assert len(by_cell) == len(audit) == len(commune_specs) == 565
    assert {spec.geography_id for spec in commune_specs} == set(statbel)
    for spec in commune_specs:
        row_number = int(re.search(r"[0-9]+$", spec.selector.address).group())
        publisher_name = nl_rows[row_number - 1][1]
        entry = by_cell[spec.selector.address]
        canonical, statbel_nl, statbel_fr = statbel[spec.geography_id]
        assert entry["publisher_name"] == publisher_name
        assert entry["geography_id"] == spec.geography_id
        assert spec.filters["pensionstat.gemeente"] == publisher_name
        assert spec.geography_level == "commune"
        assert spec.geography_vintage == "nis_2025"
        assert spec.geography_name == canonical
        fr_row = int(entry["publisher_fr_cell"][1:])
        fr_name = fr_rows[fr_row - 1][1]
        assert fr_rows[fr_row - 1][0] == 2025
        assert fr_name == entry["publisher_fr_name"]
        assert fr_series[fr_name] == nl_series[publisher_name]
        if entry["match"] == "statbel_nl_exact":
            assert publisher_name == statbel_nl
        elif entry["match"] == "statbel_nl_orthographic":
            assert publisher_name != statbel_nl
            assert _orthography(publisher_name) == _orthography(statbel_nl)
        elif entry["match"].startswith("publisher_fr_label_"):
            assert _orthography(publisher_name) != _orthography(statbel_nl)
            assert _orthography(fr_name) == _orthography(statbel_fr)
        else:
            assert entry["match"] == "publisher_section_label"
            assert (publisher_name, fr_name, spec.geography_id) == (
                "Montegnée",
                "Montegnée",
                "62093",
            )
    assert Counter(entry["match"] for entry in audit) == {
        "statbel_nl_exact": 551,
        "statbel_nl_orthographic": 9,
        "publisher_fr_label_exact": 3,
        "publisher_fr_label_orthographic": 1,
        "publisher_section_label": 1,
    }
    # The French edition labels Sint-Niklaas "Saint-Nicolas", the official
    # French name of NIS 62093, which is why 62093 carries its section name.
    sint_niklaas = next(e for e in audit if e["publisher_name"] == "Sint-Niklaas")
    assert (sint_niklaas["geography_id"], sint_niklaas["publisher_fr_name"]) == (
        "46021",
        "Saint-Nicolas",
    )
    assert statbel["62093"][2] == "Saint-Nicolas"
