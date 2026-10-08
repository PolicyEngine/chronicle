"""Author issue #322 selectors from the preserved ONS and DESNZ workbooks.

This writes declarative source packages; it does not build or validate facts.
The ONS extension retains the existing seven series' source-record identifiers.
"""

from __future__ import annotations

import re
from calendar import monthrange
from copy import deepcopy
from pathlib import Path

import openpyxl
import yaml
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
ONS_DIRECTORY = ROOT / "packages/ons/consumer_trends_current_price_2026"
ROAD_DIRECTORY = ROOT / "packages/desnz/road_transport_fuel_consumption_2024"
ONS_SOURCE = (
    "https://www.ons.gov.uk/economy/nationalaccounts/satelliteaccounts/"
    "datasets/consumertrendscurrentpricenotseasonallyadjusted"
)
ROAD_SOURCE = (
    "https://www.gov.uk/government/statistics/uk-road-transport-fuel-consumption-"
    "at-regional-and-local-authority-level-2005-to-2024"
)
ROAD_METHODOLOGY = (
    f"{ROAD_SOURCE}/sub-national-road-transport-fuel-consumption-statistics-"
    "methodology-summary-2005-2024-web-accessible"
)
ONS_SHEETS = ["0CN", *(f"{division:02}CN" for division in range(1, 13)), "TOURCN"]
ONS_BASIS = {
    "price_basis": "current_prices",
    "seasonal_adjustment": "not_seasonally_adjusted",
    "published_unit": "gbp_million",
}


class MergeKey(str):
    """A YAML merge key, distinct from an ordinary quoted string key."""


class PackageDumper(yaml.SafeDumper):
    """Emit reusable selector templates as standard YAML merges."""


PackageDumper.add_representer(
    MergeKey,
    lambda dumper, _key: dumper.represent_scalar("tag:yaml.org,2002:merge", "<<"),
)


def constraints(filters: dict) -> list[dict]:
    """Keep publisher classifications queryable as first-class constraints."""
    return [
        {
            "variable": key,
            "operator": "==",
            "value": value,
            "label": key.replace("_", " ").capitalize(),
        }
        for key, value in filters.items()
    ]


def period(label: int | str | None) -> tuple | None:
    """Represent a publisher calendar year or quarter without aligning it."""
    if isinstance(label, int) and 2020 <= label <= 2026:
        return (
            "calendar_year",
            label,
            f"cy{label}",
            {
                "start_date": f"{label}-01-01",
                "end_date": f"{label}-12-31",
                "basis": "calendar",
                "source_period_label": str(label),
            },
            {"frequency": "annual"},
        )
    match = re.fullmatch(r"(202[0-6]) Q([1-4])", str(label))
    if match:
        year, quarter = map(int, match.groups())
        first, last = 3 * quarter - 2, 3 * quarter
        return (
            "quarter",
            f"{year}-Q{quarter}",
            f"{year}_q{quarter}",
            {
                "start_date": f"{year}-{first:02}-01",
                "end_date": f"{year}-{last:02}-{monthrange(year, last)[1]:02}",
                "basis": "calendar",
                "source_period_label": str(label),
            },
            {"frequency": "quarterly", "quarter": f"Q{quarter}"},
        )
    return None


def coicop_code(raw: float | str) -> str:
    """Preserve ONS identifiers, including Excel's numeric 10, 11 and 12."""
    return str(raw).zfill(2) if isinstance(raw, int) and raw != 0 else str(raw)


def write_package(directory: Path, payload: dict, comment: str) -> None:
    """Write compact YAML with category templates shared across periods."""
    directory.mkdir(parents=True, exist_ok=True)
    (directory / "source_package.yaml").write_text(
        comment
        + yaml.dump(payload, Dumper=PackageDumper, sort_keys=False, allow_unicode=True),
        encoding="utf-8",
    )


def build_ons() -> None:
    """Extend the existing pinned ONS package without downloading a new vintage."""
    path = ONS_DIRECTORY / "source_package.yaml"
    payload = yaml.safe_load(path.read_text())
    legacy = [
        record
        for record in payload["record_sets"]
        if not record["record_set_id"].startswith("ons.consumer_trends.expanded_")
    ]
    payload["record_sets"] = legacy
    payload["label"] = (
        "ONS Consumer Trends household final consumption expenditure, every COICOP "
        "division and class, national/domestic totals and tourism, 2020 to 2026 Q1"
    )
    payload["artifact"].update(
        {
            "source_table": "Consumer Trends current price, not seasonally adjusted, 0CN, 01CN-12CN and TOURCN",
            "sheets": ["Notes", *ONS_SHEETS],
            "extraction_method": "XLSX used-range cell parse of declared sheets",
        }
    )
    payload["dimension_labels"].update(
        {
            "source_sheet": "Publisher worksheet",
            "price_basis": "Price basis",
            "seasonal_adjustment": "Seasonal adjustment",
            "published_unit": "Published unit",
            "consumption_concept": "Consumption concept",
            "cdid": "ONS series identifier",
            "publication_status": "Publication status",
            "tourism_flow": "Tourism flow",
        }
    )
    selected = {
        (record["sheet_name"], record["rows"][0]["filters"]["coicop"])
        for record in legacy
    }
    workbook = openpyxl.load_workbook(
        ROOT / "db/data/ons/consumer_trends_current_price_2026/cpnsa.xlsx",
        read_only=True,
        data_only=True,
    )
    try:
        tables = {
            name: list(workbook[name].iter_rows(values_only=True))
            for name in ONS_SHEETS
        }
    finally:
        workbook.close()
    for record in legacy:
        sheet = record["sheet_name"]
        record["entity_role"] = "households_on_uk_territory"
        shared = {**ONS_BASIS, "source_sheet": sheet, "consumption_concept": "domestic"}
        record["shared_filters"] = shared
        record["shared_constraints"] = constraints(shared)
        measure = record["measures"][0]
        column = openpyxl.utils.column_index_from_string(measure["column"])
        header = 7 if record["period_type"] == "calendar_year" else 41
        measure["expected_column_header_row"] = header
        measure["expected_column_header"] = tables[sheet][header - 1][column - 1]
        record["rows"][0]["guard_cells"] = [
            {
                "column": measure["column"],
                "row": header + 1,
                "expected_value": tables[sheet][header][column - 1],
                "label": "ONS CDID",
            }
        ]

    numeric_measure = [
        {
            "measure_id": "expenditure",
            "label": "Household final consumption expenditure",
            "ordinal": 0,
            "column": "B",
            "concept": "ons.household_final_consumption_expenditure",
            "source_concept": "ons.household_final_consumption_expenditure",
            "concept_relation": "source_label",
            "concept_authority": "ons",
            "concept_evidence_url": ONS_SOURCE,
            "concept_evidence_notes": "Published HFCE series; COICOP, CDID and domestic/national/tourism basis remain dimensions. Publisher £ million is represented in GBP; signed tourism adjustments are retained.",
            "unit": "gbp",
            "aggregation": "sum",
            "expected_cell_type": "number",
            "value_scale": 1_000_000,
        }
    ]
    marked_measure = deepcopy(numeric_measure)
    marked_measure[0].update({"expected_cell_type": "text", "value_scale": 1})
    # Reuse category declarations across periods so YAML stays compact.
    categories = {}
    for sheet, table in tables.items():
        for column, raw_code in enumerate(table[6][1:], 2):
            code = coicop_code(raw_code)
            label = str(table[5][column - 1]).replace("\n", " ")
            payload["dimension_value_labels"]["coicop"].setdefault(
                code, f"{code} {label}"
            )
            classification = {
                "coicop": code,
                "cdid": table[7][column - 1],
                "consumption_concept": "national"
                if code == "NAT0"
                else "national_adjustment"
                if code.startswith("TOUR")
                else "domestic",
            }
            if code.startswith("TOUR"):
                classification["tourism_flow"] = {
                    "TOUR": "net_tourism",
                    "TOUR1": "non_residents_spending_in_uk",
                    "TOUR2": "residents_spending_abroad",
                }[code]
            categories[sheet, column] = classification, constraints(classification)

    row_templates = {}
    for sheet, table in tables.items():
        for row_number, values in enumerate(table, 1):
            parsed = period(values[0])
            if parsed is None or (sheet == "TOURCN" and parsed[0] != "calendar_year"):
                continue
            period_type, value, suffix, coverage, frequency = parsed
            header = 7 if period_type == "calendar_year" else 41
            shared = {**ONS_BASIS, "source_sheet": sheet, **frequency}
            groups = {"published": [], "not_available": []}
            for column, raw_code in enumerate(table[6][1:], 2):
                code = coicop_code(raw_code)
                if (sheet, code) in selected:
                    continue
                raw_value = values[column - 1]
                if not isinstance(raw_value, (int, float)) and raw_value != "[x]":
                    raise ValueError(
                        f"Unexpected ONS cell {sheet}!{get_column_letter(column)}{row_number}: {raw_value!r}"
                    )
                status = "not_available" if raw_value == "[x]" else "published"
                classification, category_constraints = categories[sheet, column]
                template_key = (sheet, column, header)
                if template_key not in row_templates:
                    row_templates[template_key] = {
                        "value_id": f"coicop_{code.lower().replace('.', '_')}",
                        "label": str(table[5][column - 1]).replace("\n", " "),
                        "ordinal": column - 2,
                        "column": get_column_letter(column),
                        "source_column_id": classification["cdid"],
                        "expected_row_header_column": "A",
                        "expected_column_header_row": header,
                        "expected_column_header": table[header - 1][column - 1],
                        "guard_cells": [
                            {
                                "column": get_column_letter(column),
                                "row": header + 1,
                                "expected_value": table[header][column - 1],
                                "label": "ONS CDID",
                            }
                        ],
                        "table_record_kind": "detail",
                        "filters": classification,
                        "constraints": category_constraints,
                    }
                groups[status].append(
                    {
                        MergeKey("<<"): row_templates[template_key],
                        "row_number": row_number,
                        "expected_row_header": values[0],
                    }
                )
            for status, rows in groups.items():
                if not rows:
                    continue
                identifier = (
                    f"ons.consumer_trends.expanded_{sheet.lower()}.{suffix}.{status}"
                )
                filters = {**shared, "publication_status": status}
                payload["record_sets"].append(
                    {
                        "record_set_id": identifier,
                        "provenance_class": "administrative",
                        "assertion": "observation",
                        "record_set_spec_id": f"{identifier}.v1",
                        "source_record_id_prefix": identifier,
                        "sheet_name": sheet,
                        "period_type": period_type,
                        "period": value,
                        "period_coverage": coverage,
                        "geography_id": "K02000001",
                        "geography_level": "country",
                        "geography_name": "United Kingdom",
                        "geography_vintage": "current",
                        "entity": "household",
                        "entity_role": "household_consumers",
                        "domain": "household_consumption_expenditure",
                        "groupby_dimension": "ons.coicop",
                        "shared_filters": filters,
                        "shared_constraints": constraints(filters),
                        "rows": rows,
                        "measures": marked_measure
                        if status == "not_available"
                        else numeric_measure,
                    }
                )
    write_package(
        ONS_DIRECTORY,
        payload,
        "# Same preserved September 2026 workbook; no new download.\n"
        "# Published [x] markers and signed tourism adjustments are retained.\n"
        "# Regenerate with scripts/build_uk_consumption_tax_packages.py.\n",
    )


def build_road_fuel() -> None:
    """Select publisher national, country and regional vehicle/fuel totals."""
    workbook = openpyxl.load_workbook(
        ROOT / "db/data/desnz/road_transport_fuel_consumption_2024/"
        "subnational-road-transport-fuel-consumption-tables-2005-2024.xlsx",
        read_only=True,
        data_only=True,
    )
    vehicle_columns = [
        ("G", "buses_and_coaches", "diesel"),
        ("K", "cars", "diesel"),
        ("O", "cars", "petrol"),
        ("S", "motorcycles", "petrol"),
        ("W", "hgv", "diesel"),
        ("AA", "hgv", "natural_gas"),
        ("AE", "lgv", "diesel"),
        ("AI", "lgv", "petrol"),
        ("AM", "lgv", "lpg"),
        ("AP", "all_vehicles", "all_fuels"),
    ]
    basis = {
        "population_scope": "vehicles_travelling_in_area",
        "electric_vehicle_treatment": "excluded",
        "estimate_basis": "spatial_estimates_without_dukes_normalisation",
        "national_inventory_basis": "normalised_to_adjusted_dukes_fuel_sales",
        "road_type": "all_roads",
    }
    try:
        tables = {
            year: list(workbook[str(year)].iter_rows(values_only=True))
            for year in range(2005, 2025)
        }
    finally:
        workbook.close()
    measures = []
    for ordinal, (column, vehicle, fuel) in enumerate(vehicle_columns):
        filters = {"vehicle_type": vehicle, "fuel": fuel}
        if fuel in {"petrol", "diesel"}:
            filters["biofuel_treatment"] = "includes_blended_biofuels"
        label = tables[2024][3][openpyxl.utils.column_index_from_string(column) - 1]
        measures.append(
            {
                "measure_id": f"{vehicle}_{fuel}",
                "label": label,
                "ordinal": ordinal,
                "column": column,
                "source_column_id": label,
                "concept": "desnz.road_transport.fuel_consumption",
                "source_concept": "desnz.road_transport.fuel_consumption",
                "concept_relation": "source_label",
                "concept_authority": "desnz",
                "concept_evidence_url": ROAD_METHODOLOGY,
                "concept_evidence_notes": "Fuel used by vehicles travelling in the area, not vehicles registered to residents. Spatial mapping is not normalised to fuel sales; the national emissions inventory is normalised to adjusted DUKES totals. Electric-vehicle kilometres are removed. Workbook Note 3 supplies the petrol/diesel classification and Note 6 includes blended biofuels. No ktoe-to-litres or ktoe-to-tonnes factors are published here.",
                "unit": "ktoe",
                "aggregation": "sum",
                "expected_cell_type": "number",
                "expected_column_header_row": 4,
                "expected_column_header": label,
                "filters": filters,
                "constraints": constraints(filters),
            }
        )
    records = []
    for year, table in tables.items():
        rows = []
        for row_number, values in enumerate(table, 1):
            code, region, name = values[:3]
            if isinstance(code, str) and code.startswith(
                ("K0", "E12", "E13", "W92", "S92", "N92")
            ):
                geography = code
                header_column, expected = "A", code
            elif code is None and region == name == "England":
                geography = "E92000001"
                header_column, expected = "C", name
            else:
                continue
            rows.append(
                {
                    "value_id": geography.lower(),
                    "label": name,
                    "ordinal": len(rows),
                    "row_number": row_number,
                    "table_record_kind": "total",
                    "expected_row_header_column": header_column,
                    "expected_row_header": expected,
                    "guard_cells": [
                        {
                            "column": "C",
                            "row": "start",
                            "expected_value": name,
                            "label": "Published area name",
                        },
                        {
                            "column": "A",
                            "row": 1,
                            "expected_value": table[0][0],
                            "label": "Calendar year and ktoe unit",
                        },
                    ],
                    "geography_id": geography,
                    "geography_level": "region"
                    if geography.startswith(("E12", "E13"))
                    else "country",
                    "geography_name": str(name).removesuffix(" total"),
                    "geography_vintage": "2025-12",
                }
            )
        identifier = f"desnz.road_transport.fuel_consumption.cy{year}"
        records.append(
            {
                "record_set_id": identifier,
                "provenance_class": "administrative",
                "assertion": "observation",
                "record_set_spec_id": f"{identifier}.v1",
                "source_record_id_prefix": identifier,
                "sheet_name": str(year),
                "period_type": "calendar_year",
                "period": year,
                "period_coverage": {
                    "start_date": f"{year}-01-01",
                    "end_date": f"{year}-12-31",
                    "basis": "calendar",
                    "source_period_label": str(year),
                },
                "geography_id": "K02000001",
                "geography_level": "country",
                "geography_name": "United Kingdom",
                "geography_vintage": "2025-12",
                "entity": "institutional_sector",
                "entity_role": "road_transport",
                "domain": "road_transport_fuel_consumption",
                "groupby_dimension": "desnz.road_transport.area",
                "shared_filters": basis,
                "shared_constraints": constraints(basis),
                "rows": rows,
                "measures": measures,
            }
        )
    payload = {
        "schema_version": "ledger.source_package.v1",
        "package_id": "desnz-road-transport-fuel-consumption-2024",
        "label": "DESNZ road transport fuel consumption by vehicle type and fuel, UK, GB, countries and regions, 2005-2024",
        "dimension_labels": {
            "desnz.road_transport.area": "Published geography",
            **{
                key: key.replace("_", " ").capitalize()
                for key in [*basis, "vehicle_type", "fuel", "biofuel_treatment"]
            },
        },
        "artifact": {
            "source_name": "desnz",
            "source_table": "Sub-national road transport fuel consumption statistics, 2005-2024",
            "resource_package": "db",
            "resource_directory": "data/desnz/road_transport_fuel_consumption_2024",
            "manifest": "manifest.yaml",
            "vintage": "2026_06_25_release_2005_2024",
            "extracted_at": "2026-10-08",
            "extraction_method": "XLSX used-range cell parse of declared sheets",
            "parser": "xlsx_used_range",
            "artifact_year": 2026,
            "sheets": ["Contents", "Notes", *(str(year) for year in range(2005, 2025))],
        },
        "record_sets": records,
    }
    write_package(
        ROAD_DIRECTORY,
        payload,
        "# Published vehicle/fuel totals only; no conversion or reconciliation.\n"
        "# Local-authority cells remain in the preserved workbook, outside emitted facts.\n"
        "# Regenerate with scripts/build_uk_consumption_tax_packages.py.\n",
    )


if __name__ == "__main__":
    build_ons()
    build_road_fuel()
