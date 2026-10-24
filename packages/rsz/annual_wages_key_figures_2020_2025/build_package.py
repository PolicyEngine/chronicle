#!/usr/bin/env python3
"""Transcribe RSZ annual key figures at the workbook's displayed precision.

One CSV row is one published XLSX data cell. Cached numeric values are rounded
only to the cell's declared number format; published thousands are never scaled.
The original workbook, labels, coordinates, cached values and format survive.
"""

from __future__ import annotations

import csv
from decimal import Decimal, ROUND_HALF_UP
import hashlib
from pathlib import Path
import re

from openpyxl import load_workbook
import yaml

PACKAGE_DIR = Path(__file__).resolve().parent
REPO_ROOT = PACKAGE_DIR.parents[2]
DATA_DIR = REPO_ROOT / "db/data/rsz/annual_wages_key_figures_2020_2025"
PACKAGE_ID = "rsz-annual-wages-key-figures-2020-2025"
WORKBOOK = "rsz-ruwedata-periods-wages.xlsx"
SHEET = "Kerncijfers-Chiffres Clés"
CSV_NAME = "annual_wages_key_figures_2020_2025.csv"
WORKBOOK_SHA256 = "7eea6aa74dd8c3c287c54bf591fef0e8e5d5507404988f94a05bf259fd6883af"
CAPTURED_AT = "2026-10-24T01:05:27Z"
LANDING_URL = "https://www.rsz.be/stats/voorpublicatie-ruwe-data"
WORKBOOK_URL = (
    "https://www.rsz.be/file/6103c41969844d308979992d7f84d51e59927e1e/"
    "c4c5776181e3dd24273ecb6b9486065cb468fadb/"
    "ruwedata-donneesbrutes-periods-wages.xlsx"
)
SOURCE_TABLE = "Lonen en bezoldigde prestaties, Kerncijfers-Chiffres Clés, 2020–2025"
DIMENSIONS = {
    "rsz.annual_worker_type": "werknemerstype",
    "rsz.annual_sex": "Geslacht / Sexe",
    "rsz.annual_worker_type_sex": "werknemerstype / Geslacht / Sexe",
}
METRICS = (
    ("D", "ordinary_remuneration", "EUR_thousands"),
    ("E", "severance", "EUR_thousands"),
    ("F", "premiums", "EUR_thousands"),
    ("G", "simple_departure_holiday_pay", "EUR_thousands"),
    ("H", "full_time_paid_days", "day"),
    ("I", "part_time_paid_hours", "hour_thousands"),
)
EXPECTED_HEADERS = {
    "D": " gewone bezoldiging \nrémunération ordinaire",
    "E": " verbrekingsvergoeding\nindemnité de rupture",
    "F": " premies \nprimes",
    "G": " enkel vertrekvakantiegeld \npécule simple de sortie",
    "H": " Dagen Voltijds \n Journées Temps Plein",
    "I": " Uren  deeltijds \n Heures temps partiel\n(x 1.000)",
}
WORKER_IDS = {
    "Arbeider - Ouvrier": "arbeider",
    "Bediende - Employé": "bediende",
    "Ambtenaar - Fonctionnaire": "ambtenaar",
}
SEX_IDS = {"Man - Homme": "man", "Vrouw - Femme": "vrouw"}
HEADERS = [
    "row_id",
    "year",
    "workbook",
    "sheet",
    "cell",
    "source_row",
    "source_column",
    "year_printed",
    "worker_type_printed",
    "sex_printed",
    "measure_printed",
    "number_format",
    "raw_numeric_value",
    *DIMENSIONS,
    "measure_id",
    "value",
]
EVIDENCE = {
    "source_capture": (WORKBOOK, WORKBOOK_URL, WORKBOOK_SHA256),
    "landing_page": (
        "rsz-voorpublicatie-ruwe-data.html",
        LANDING_URL,
        "5fff81ac54d2661573e2b8b736ab82d900ea1a5b89c60b591ec1cf5824560ddb",
    ),
    "daily_wage_definition": (
        "rsz-gemeenschappelijke-criteria.html",
        "https://www.rsz.be/gemeenschappelijke-criteria",
        "b1f9df57fcd06e5091ed34d642ddb91a4c79b3d86f1d50cbaf68fc3e41b54c91",
    ),
    "variable_definitions": (
        "rsz-gemeenschappelijke-variabelen.html",
        "https://www.rsz.be/gemeenschappelijke-variabelen",
        "e008e4a37f1686b8ce8c5bd22b575e3953b76dbfb701e7490f2ce6ec099f0428",
    ),
    "licence_evidence": (
        "rsz-open-data-data-governance-act.html",
        "https://www.rsz.be/open-data-data-governance-act",
        "956cb55a26adbdca779ed153a8cb8290bb2dea3e2b643380cc485fc5d554ab6e",
    ),
    "capture_receipt": ("CAPTURED_AT.txt", LANDING_URL, None),
    "capture_checksums": ("SHA256SUMS.txt", LANDING_URL, None),
}
SCOPE = (
    "RSZ-registered employment in Belgium; includes inbound cross-border workers, "
    "excludes posted workers and self-employed people. Administrative employment "
    "population differs from Statbel fiscal declarations."
)


class LiteralDumper(yaml.SafeDumper):
    def ignore_aliases(self, data):
        return True


def yaml_text(value: dict) -> str:
    return yaml.dump(
        value, Dumper=LiteralDumper, sort_keys=False, allow_unicode=True, width=100
    )


def displayed_decimal(cell) -> str:
    """Read mandatory decimal places from this publisher's fixed-point format."""
    number_format = cell.number_format.split(";", 1)[0]
    match = re.fullmatch(r"[#,0]+(?:\.(0+))?", number_format)
    if not match:
        raise ValueError(
            f"Unsupported number format {cell.coordinate}: {number_format!r}"
        )
    if isinstance(cell.value, bool) or not isinstance(cell.value, (int, float)):
        raise ValueError(f"Not a numeric data cell: {cell.coordinate}={cell.value!r}")
    places = len(match[1] or "")
    quantum = Decimal(1).scaleb(-places)
    return format(
        Decimal(str(cell.value)).quantize(quantum, rounding=ROUND_HALF_UP), "f"
    )


def extract_rows(data_dir: Path = DATA_DIR) -> list[dict]:
    raw_path = data_dir / WORKBOOK
    if hashlib.sha256(raw_path.read_bytes()).hexdigest() != WORKBOOK_SHA256:
        raise ValueError("RSZ annual wages workbook checksum mismatch")
    values = load_workbook(raw_path, data_only=True, read_only=True)
    formulas = load_workbook(raw_path, data_only=False, read_only=True)
    rows = []
    try:
        sheet, formula_sheet = values[SHEET], formulas[SHEET]
        if (sheet.max_row, sheet.max_column) != (45, 15):
            raise ValueError("RSZ annual key-figures used range changed")
        if sheet["C2"].value != "Bezoldiging - rémunération (x 1.000 Eur)":
            raise ValueError("RSZ monetary unit heading changed")
        for column, expected in EXPECTED_HEADERS.items():
            if sheet[f"{column}3"].value != expected:
                raise ValueError(f"RSZ measure heading changed at {column}3")
        # Read entire tiny sheet once: repeated read-only cell lookups reparse XML.
        cached_rows = list(sheet.iter_rows())
        formula_rows = list(formula_sheet.iter_rows())
        for year in range(2020, 2026):
            start = 4 + (year - 2020) * 7
            worker = None
            for source_row in range(start, start + 7):
                cached = cached_rows[source_row - 1]
                authored = formula_rows[source_row - 1]
                printed_year = cached[0].value
                total = source_row == start + 6
                if total:
                    if printed_year != f"Totaal {year}":
                        raise ValueError(f"Unexpected total label at A{source_row}")
                    worker_label = sex = ""
                    native, label = "total", "Totaal"
                else:
                    if source_row == start and printed_year != year:
                        raise ValueError(f"Unexpected year at A{source_row}")
                    if cached[1].value is not None:
                        worker = cached[1].value
                    worker_label, sex = worker, cached[2].value
                    native = f"{WORKER_IDS[worker_label]}.{SEX_IDS[sex]}"
                    label = f"{worker_label} / {sex}"
                for column, metric, _ in METRICS:
                    index = ord(column) - ord("A")
                    formula = authored[index].value
                    if authored[index].data_type == "f":
                        if total and str(formula).upper().startswith("=SUBTOTAL("):
                            continue  # Validation totals never become facts.
                        raise ValueError(
                            f"Unexpected formula at {SHEET}!{column}{source_row}"
                        )
                    cell = cached[index]
                    row = dict.fromkeys(HEADERS, "")
                    row.update(
                        row_id=f"{year}.{native}.{metric}",
                        year=str(year),
                        workbook=WORKBOOK,
                        sheet=SHEET,
                        cell=cell.coordinate,
                        source_row=str(source_row),
                        source_column=column,
                        year_printed="" if printed_year is None else str(printed_year),
                        worker_type_printed=cached[1].value or "",
                        sex_printed=cached[2].value or "",
                        measure_printed=EXPECTED_HEADERS[column],
                        number_format=cell.number_format,
                        raw_numeric_value=str(cell.value),
                        measure_id=metric,
                        value=displayed_decimal(cell),
                    )
                    row["rsz.annual_worker_type"] = worker_label
                    row["rsz.annual_sex"] = sex
                    row["rsz.annual_worker_type_sex"] = native
                    row["_label"] = label
                    rows.append(row)
    finally:
        values.close()
        formulas.close()
    return rows


def package_spec(rows: list[dict]) -> dict:
    column_by_field = {name: chr(ord("A") + i) for i, name in enumerate(HEADERS)}
    labels = {
        key: {row[key]: row[key] for row in rows if row[key]} for key in DIMENSIONS
    }
    labels["rsz.annual_worker_type_sex"] = {
        row["rsz.annual_worker_type_sex"]: row["_label"] for row in rows
    }
    record_sets = []
    for year in range(2020, 2026):
        identifier = f"rsz.annual_wages.year{year}"
        for column, metric, unit in METRICS:
            members = [
                (csv_row, row)
                for csv_row, row in enumerate(rows, 2)
                if row["year"] == str(year) and row["measure_id"] == metric
            ]
            package_rows = []
            for ordinal, (csv_row, row) in enumerate(members):
                native = row["rsz.annual_worker_type_sex"]
                filters = {key: row[key] for key in DIMENSIONS if row[key]}
                guards = [
                    {
                        "column": column_by_field[key],
                        "expected_value": int(row[key])
                        if key == "year"
                        else row[key].strip(),
                        "label": key,
                    }
                    for key in (
                        "year",
                        "workbook",
                        "sheet",
                        "cell",
                        "measure_id",
                        "number_format",
                        "measure_printed",
                    )
                ]
                package_rows.append(
                    {
                        "value_id": native,
                        "label": row["_label"],
                        "ordinal": ordinal,
                        "row_number": csv_row,
                        "source_row_id": f"{year}.{native}",
                        "expected_row_header_column": "A",
                        "expected_row_header": row["row_id"],
                        "table_record_kind": "total" if native == "total" else "detail",
                        "filters": filters,
                        "constraints": [
                            {
                                "variable": key,
                                "operator": "==",
                                "value": value,
                                "label": DIMENSIONS[key],
                            }
                            for key, value in filters.items()
                        ],
                        "guard_cells": guards,
                    }
                )
            record_sets.append(
                {
                    "record_set_id": f"{identifier}.{metric}",
                    "provenance_class": "administrative",
                    "assertion": "observation",
                    # Shared table spec and record prefix keep IDs independent
                    # of a later measure split or coverage clarification.
                    "record_set_spec_id": "rsz.annual_wages.key_figures.v1",
                    "source_record_id_prefix": identifier,
                    "sheet_name": "annual_wages_key_figures",
                    "period_type": "calendar_year",
                    "period": year,
                    "period_coverage": {
                        "start_date": f"{year}-01-01",
                        "end_date": f"{year}-12-31",
                        "basis": "calendar",
                        "source_period_label": str(year),
                        "notes": (
                            "Calendar-year remuneration or paid days/hours; pre-publication "
                            "raw data, landing page updated 22/10/2026. The 2024–2025 figures "
                            "are not described as final. " + SCOPE
                        ),
                    },
                    "geography_id": "BE",
                    "geography_level": "country",
                    "geography_name": "Belgium",
                    "geography_vintage": "current",
                    "entity": "person",
                    "entity_role": "employee",
                    "domain": "employment",
                    "groupby_dimension": "rsz.annual_worker_type_sex",
                    "rows": package_rows,
                    "measures": [
                        {
                            "measure_id": metric,
                            "label": EXPECTED_HEADERS[column].strip(),
                            "ordinal": 0,
                            "column": column_by_field["value"],
                            "source_column_id": metric,
                            "expected_column_header_row": 1,
                            "expected_column_header": "value",
                            "concept": f"rsz.annual_wages.{metric}",
                            "source_concept": f"rsz.annual_wages.{metric}",
                            "concept_relation": "source_label",
                            "concept_authority": "RSZ/ONSS Directie Statistiek",
                            "concept_evidence_url": WORKBOOK_URL,
                            "concept_evidence_notes": (
                                f"One published {WORKBOOK}, {SHEET}, column {column} data cell; "
                                "exact workbook/sheet/cell locator is retained and guarded in "
                                "the extract. Values retain the cell's displayed precision. "
                                f"Published unit {unit}; no rescaling, computed totals, "
                                "ratios, averages or period alignment."
                            ),
                            "unit": unit,
                            "aggregation": "sum",
                            "expected_cell_type": "number",
                        }
                    ],
                }
            )
    return {
        "schema_version": "ledger.source_package.v1",
        "package_id": PACKAGE_ID,
        "label": "RSZ/ONSS annual wages and paid work key figures, 2020–2025, October 2026 pre-publication",
        "dimension_labels": DIMENSIONS,
        "dimension_value_labels": labels,
        "artifact": {
            "source_name": "rsz_annual_wages_key_figures",
            "source_table": SOURCE_TABLE,
            "resource_package": "db",
            "resource_directory": "data/rsz/annual_wages_key_figures_2020_2025",
            "manifest": "manifest.yaml",
            "vintage": "rsz_raw_wages_prepublication_2026_10_22",
            "extracted_at": "2026-10-24",
            "extraction_method": "build_package.py; one CSV row per publisher XLSX cell; Decimal at number-format precision; no rescaling",
            "parser": "delimited_text_full_rows",
            "artifact_year": 2025,
            "sheet_name": "annual_wages_key_figures",
        },
        "record_sets": record_sets,
    }


def generate(data_dir: Path = DATA_DIR, package_dir: Path = PACKAGE_DIR) -> None:
    rows = extract_rows(data_dir)
    with (data_dir / CSV_NAME).open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(
            stream, fieldnames=HEADERS, lineterminator="\n", extrasaction="ignore"
        )
        writer.writeheader()
        writer.writerows(rows)
    entries = {2025: (CSV_NAME, WORKBOOK_URL, None), **EVIDENCE}
    files = {}
    for key, (name, url, expected_sha) in entries.items():
        path = data_dir / name
        sha = hashlib.sha256(path.read_bytes()).hexdigest()
        if expected_sha is not None and sha != expected_sha:
            raise ValueError(f"Capture checksum mismatch: {name}")
        r2_key = f"raw/belgium/{PACKAGE_ID}/{key}/{sha}/{name}"
        files[key] = {
            "filename": name,
            "source_url": url,
            "source_table": SOURCE_TABLE,
            "sha256": sha,
            "size_bytes": path.stat().st_size,
            "retrieved_at": CAPTURED_AT,
            "notes": (
                "Representation-only one-cell transcription; exact workbook, sheet, cell, "
                "cached value, displayed precision and printed labels are retained. "
                if key == 2025
                else "Unmodified staged source capture or acquisition evidence. "
            )
            + "Uploaded to R2 ledger-raw and verified by SHA-256 readback on 2026-10-24.",
            "storage": {
                "r2": {
                    "provider": "r2",
                    "bucket": "ledger-raw",
                    "key": r2_key,
                    "uri": f"r2://ledger-raw/{r2_key}",
                }
            },
        }
    (data_dir / "manifest.yaml").write_text(
        yaml_text(
            {
                "source_id": "belgium",
                "package_id": PACKAGE_ID,
                "source_name": "RSZ/ONSS annual remuneration and paid work key figures",
                "publisher": "Rijksdienst voor Sociale Zekerheid / Office national de sécurité sociale, Directie Statistiek",
                "source_page": LANDING_URL,
                "license": "CC BY 4.0",
                "license_url": "https://creativecommons.org/licenses/by/4.0/",
                "license_evidence_url": EVIDENCE["licence_evidence"][1],
                "attribution": "RSZ/ONSS, Directie Statistiek, Lonen en bezoldigde prestaties, October 2026 pre-publication; retrieved 2026-10-24. Representation normalized by PolicyEngine Chronicle; CC BY 4.0.",
                "files": files,
            }
        ),
        encoding="utf-8",
    )
    package_dir.mkdir(parents=True, exist_ok=True)
    (package_dir / "source_package.yaml").write_text(
        yaml_text(package_spec(rows)), encoding="utf-8"
    )
    typed_totals = sum(row["rsz.annual_worker_type_sex"] == "total" for row in rows)
    print(
        f"{len(rows)} published cells: {len(rows) - typed_totals} detail, {typed_totals} typed total cells"
    )


if __name__ == "__main__":
    generate()
