#!/usr/bin/env python3
"""Select verbatim 2025 PensionStat cells; never aggregate or rescale values."""

from __future__ import annotations

from collections import defaultdict
import csv
import hashlib
import json
from pathlib import Path
import unicodedata

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import yaml


PACKAGE_DIR = Path(__file__).resolve().parent
REPO_ROOT = PACKAGE_DIR.parents[2]
DATA_DIR = REPO_ROOT / "db/data/sfpd/legal_pension_caseload_2025"
RAW_PATH = DATA_DIR / "data-2025-nl.xlsx"
SHA256 = "d358fff0c643c8edfd53b55514c311d7f98a03e3c7d5a37be1450e5ce84f94a2"
SOURCE_URL = (
    "https://www.pensionstat.be/file/6103c41969844d308979992d7f84d51e59927e1e/"
    "b7bc512baaae29d4344ac44519e892339ddf5d19/data-2025-nl.xlsx"
)
SOURCE_PAGE = (
    "https://www.pensionstat.be/nl/kerncijfers/wettelijk-pensioen/gepensioneerden"
)
FR_RAW_PATH = DATA_DIR / "data-2025-fr.xlsx"
FR_SHA256 = "af638bd6dc344e83b514d91c7c2cb0b5b4cd00a8da10f8659ac452fbe8b79c33"
STATBEL_COMMUNES = (
    REPO_ROOT
    / "db/data/statbel/population_structure_commune_2025"
    / "statbel_population_structure_commune_2025.csv"
)
# PensionStat labels Saint-Nicolas (Liege province, NIS 62093) by its section
# Montegnee in both editions: its French edition gives the label
# "Saint-Nicolas" to Sint-Niklaas (NIS 46021).
SECTION_LABELS = {"Montegnée": "62093"}
FR_LABEL_COLLISIONS = {"Sint-Niklaas": "Saint-Nicolas"}
PACKAGE_ID = "sfpd-legal-pension-caseload-2025"
SOURCE_TABLE = "PensionStat legal pensions and IGO, 2025"
SHEETS = {
    "Gepensioneerden": (
        "Jaar",
        "Pensioentype",
        "Geslacht",
        "Burgerlijke staat",
        "Nationaliteit",
        "Woonplaats",
        "Leeftijdsklasse",
        "Stelsel",
        "Aantal",
        "Gemiddeld bedrag januari",
    ),
    "Gepensioneerden_gemeente": ("Jaar", "Gemeente", "Aantal"),
    "IGO": (
        "Jaar",
        "Pensioentype",
        "Geslacht",
        "Nationaliteit",
        "Regio",
        "Burgerlijke staat",
        "Leeftijdsklasse",
        "Verhoogd basisbedrag",
        "Aantal",
        "Gemiddeld bedrag",
        "Mediaanbedrag",
    ),
    "Jaaruitgave": (
        "Jaar",
        "Uit het kadaster",
        "Stelsel",
        "Type",
        "Berekeningstarief",
        "Type pensioen of uitkering",
        "Jaaruitgave",
    ),
    "MinP_aantal": (
        "Jaar",
        "Pensioentype",
        "Geslacht",
        "Burgerlijke staat",
        "Nationaliteit",
        "Woonplaats",
        "Leeftijdsklasse",
        "Stelsel",
        "Aantal",
    ),
}
VALUE_COLUMNS = {
    "Gepensioneerden": (8, 9),
    "Gepensioneerden_gemeente": (2,),
    "IGO": (8, 9, 10),
    "Jaaruitgave": (6,),
    "MinP_aantal": (8,),
}
_KADASTER = "Pensioners paid only through the kadaster are excluded."
# The IGO page dates its counts to 1 January 2025 and its amounts to January
# 2025; the other monthly sheets date both to January 2025.
COVERAGE = {
    "annual": {
        "start_date": "2025-01-01",
        "end_date": "2025-12-31",
        "source_period_label": "2025",
        "basis": "calendar",
        "notes": "Jaaruitgave uses the publisher's Jaar label without shifting it.",
    },
    "january": {
        "start_date": "2025-01-01",
        "end_date": "2025-01-31",
        "source_period_label": "januari 2025",
        "basis": "calendar",
        "notes": "January caseload/payment reference; gross published amounts. "
        + _KADASTER,
    },
    "igo_counts": {
        "start_date": "2025-01-01",
        "end_date": "2025-01-01",
        "source_period_label": "1 januari 2025",
        "basis": "calendar",
        "notes": "IGO counts refer to 1 January 2025 (PensionStat IGO page). "
        + _KADASTER,
    },
    "igo_amounts": {
        "start_date": "2025-01-01",
        "end_date": "2025-01-31",
        "source_period_label": "januari 2025",
        "basis": "calendar",
        "notes": "Gross monthly IGO amounts for January 2025 (PensionStat IGO page). "
        + _KADASTER,
    },
}
# IGO's value columns split into record sets by period coverage.
IGO_RECORD_SETS = (
    ("counts", VALUE_COLUMNS["IGO"][:1], "igo_counts"),
    ("amounts", VALUE_COLUMNS["IGO"][1:], "igo_amounts"),
)
# Geography names follow the Statbel packages' NUTS1 names; the publisher's
# own code or label stays in the woonplaats/regio dimension.
REGIONS = {
    "BRU": ("BE1", "Brussels Capital Region"),
    "VLA": ("BE2", "Flemish Region"),
    "WAL": ("BE3", "Walloon Region"),
    "Brussels Hoofdstedelijk Gewest": ("BE1", "Brussels Capital Region"),
    "Vlaams Gewest": ("BE2", "Flemish Region"),
    "Waals Gewest": ("BE3", "Walloon Region"),
}


class Dumper(yaml.CSafeDumper):
    def ignore_aliases(self, data):
        return True


def yaml_text(payload: dict) -> str:
    return yaml.dump(payload, Dumper=Dumper, sort_keys=False, allow_unicode=True)


def dimension(header: str) -> str:
    return "pensionstat." + header.lower().replace(" ", "_")


def name_key(name: str) -> str:
    """Only orthography: no translation, fuzzy matching, or old-area merging."""
    text = name.casefold().replace("œ", "oe").replace("’", "'")
    return "".join(
        char
        for char in unicodedata.normalize("NFKD", text)
        if not unicodedata.combining(char)
    )


def commune_reference() -> dict[str, dict[str, str]]:
    """Statbel's NIS 2025 communes with their official Dutch and French names."""
    communes: dict[str, dict[str, str]] = {}
    with STATBEL_COMMUNES.open(newline="", encoding="utf-8") as stream:
        for row in csv.DictReader(stream):
            if row["geography_level"] != "commune":
                continue
            if row["geography_vintage"] != "nis_2025":
                raise ValueError("Unexpected Statbel commune vintage")
            entry = {
                "geography_name": row["geography_name"],
                "statbel_nl": row["source_municipality_nl"],
                "statbel_fr": row["source_municipality_fr"],
            }
            if communes.setdefault(row["geography_id"], entry) != entry:
                raise ValueError(
                    f"Inconsistent Statbel names for {row['geography_id']}"
                )
    if len(communes) != 565:
        raise ValueError(f"Expected 565 NIS 2025 communes, found {len(communes)}")
    return communes


def _name_index(communes: dict, field: str) -> tuple[dict, dict]:
    exact: dict[str, set[str]] = defaultdict(set)
    normalized: dict[str, set[str]] = defaultdict(set)
    for code, entry in communes.items():
        exact[entry[field]].add(code)
        normalized[name_key(entry[field])].add(code)
    return dict(exact), dict(normalized)


def _lookup(name: str, index: tuple[dict, dict]) -> tuple[str | None, str]:
    exact, normalized = index
    codes, how = exact.get(name, set()), "exact"
    if not codes:
        codes, how = normalized.get(name_key(name), set()), "orthographic"
    if len(codes) > 1:
        raise ValueError(f"Ambiguous commune name {name!r}: {sorted(codes)}")
    return (next(iter(codes)), how) if codes else (None, how)


def _commune_series(path: Path, sheet: str, headers: tuple) -> dict:
    """Each commune's 2019-2025 counts, keyed by name, with its 2025 row."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = iter(workbook[sheet].iter_rows(values_only=True))
        if tuple(next(rows)) != headers:
            raise ValueError(f"Header drift in {path.name}!{sheet}")
        series: dict[str, dict] = defaultdict(dict)
        row_2025 = {}
        for row_number, (year, name, count) in enumerate(rows, 2):
            if year in series[name]:
                raise ValueError(f"Duplicate {year} row for {name!r} in {path.name}")
            series[name][year] = count
            if year == 2025:
                row_2025[name] = row_number
    finally:
        workbook.close()
    years = range(2019, 2026)
    return {
        name: (tuple(counts.get(year) for year in years), row_2025[name])
        for name, counts in series.items()
    }


def commune_mapping() -> dict[str, dict]:
    """Map every Dutch-edition commune name to exactly one NIS 2025 code.

    Tiers, in order: Statbel's official Dutch name (exact, then orthographic);
    the publisher's own French-edition label for the same commune, found by its
    identical 2019-2025 count series, against Statbel's official French name;
    and the one declared section label. No translation or fuzzy matching. The
    result must be a bijection onto Statbel's 565 communes.
    """
    if hashlib.sha256(FR_RAW_PATH.read_bytes()).hexdigest() != FR_SHA256:
        raise ValueError("PensionStat French workbook SHA-256 mismatch")
    communes = commune_reference()
    nl_index = _name_index(communes, "statbel_nl")
    fr_index = _name_index(communes, "statbel_fr")
    nl = _commune_series(
        RAW_PATH, "Gepensioneerden_gemeente", SHEETS["Gepensioneerden_gemeente"]
    )
    fr = _commune_series(
        FR_RAW_PATH, "Pensionnés_commune", ("Année", "Commune", "Nombre")
    )
    fr_by_series: dict[tuple, list[str]] = defaultdict(list)
    for name, (series, _) in fr.items():
        fr_by_series[series].append(name)
    mapping = {}
    for name, (series, _) in nl.items():
        partners = fr_by_series.get(series, [])
        if len(partners) != 1:
            raise ValueError(f"No unique French-edition series for {name!r}")
        fr_name = partners[0]
        fr_code, fr_how = _lookup(fr_name, fr_index)
        code, how = _lookup(name, nl_index)
        match = f"statbel_nl_{how}"
        if code is None and fr_code is not None:
            code, match = fr_code, f"publisher_fr_label_{fr_how}"
        if code is None and name in SECTION_LABELS and fr_name == name:
            code, match = SECTION_LABELS[name], "publisher_section_label"
        if code is None:
            raise ValueError(f"Unresolved PensionStat commune {name!r}")
        if fr_code not in (None, code) and FR_LABEL_COLLISIONS.get(name) != fr_name:
            raise ValueError(f"{name!r}: French label {fr_name!r} disagrees")
        mapping[name] = {
            "match": match,
            "publisher_fr_name": fr_name,
            "publisher_fr_cell": f"C{fr[fr_name][1]}",
            "statbel_nl": communes[code]["statbel_nl"],
            "statbel_fr": communes[code]["statbel_fr"],
            "geography_id": code,
            "geography_level": "commune",
            "geography_name": communes[code]["geography_name"],
            "geography_vintage": "nis_2025",
        }
    codes = [entry["geography_id"] for entry in mapping.values()]
    if len(set(codes)) != len(codes) or set(codes) != set(communes):
        raise ValueError("PensionStat communes are not a bijection onto NIS 2025")
    for name, fr_name in FR_LABEL_COLLISIONS.items():
        # The French edition reuses the label of another commune here, which is
        # why that commune carries its section name instead.
        if mapping[name]["publisher_fr_name"] != fr_name:
            raise ValueError(f"Expected French label {fr_name!r} for {name!r}")
        if _lookup(fr_name, fr_index)[0] not in SECTION_LABELS.values():
            raise ValueError(f"French label {fr_name!r} no longer collides")
    return mapping


def commune_geography(name: str, mapping: dict[str, dict]) -> dict:
    return {
        key: mapping[name][key]
        for key in (
            "geography_id",
            "geography_level",
            "geography_name",
            "geography_vintage",
        )
    }


def geography(header_values: dict, mapping: dict[str, dict]) -> dict:
    if "Gemeente" in header_values:
        return commune_geography(header_values["Gemeente"], mapping)
    residence = header_values.get("Woonplaats", header_values.get("Regio"))
    if residence == "BUI":
        return {
            "geography_id": "pensionstat:residence:BUI",
            "geography_level": "statistical_scope",
            "geography_name": "BUI",
            "geography_vintage": "pensionstat_2025",
        }
    if residence is not None:
        code, name = REGIONS[residence]  # Fail closed on an unfamiliar category.
        return {
            "geography_id": code,
            "geography_level": "nuts1",
            "geography_name": name,
            "geography_vintage": "NUTS_2024",
        }
    return {}


def measure(sheet: str, header: str, index: int, ordinal: int) -> dict:
    is_count = header == "Aantal"
    aggregation = (
        "sum"
        if is_count or sheet == "Jaaruitgave"
        else "median"
        if header == "Mediaanbedrag"
        else "mean"
    )
    concept = "pensionstat." + sheet.lower() + "." + header.lower().replace(" ", "_")
    return {
        "measure_id": header.lower().replace(" ", "_"),
        "label": header,
        "ordinal": ordinal,
        "column": get_column_letter(index + 1),
        "source_column_id": header,
        "expected_column_header_row": 1,
        "expected_column_header": header,
        "concept": concept,
        "source_concept": concept,
        "concept_relation": "source_label",
        "concept_authority": "PensionStat.be",
        "concept_evidence_url": SOURCE_URL,
        "concept_evidence_notes": (
            f"Verbatim published {sheet} / {header} cell. No totals, ratios, "
            "count-times-average amounts, rounding, scaling, or period alignment."
        ),
        "unit": "count" if is_count else "EUR",
        "aggregation": aggregation,
        "expected_cell_type": "number",
    }


def generate_package() -> dict:
    if hashlib.sha256(RAW_PATH.read_bytes()).hexdigest() != SHA256:
        raise ValueError("PensionStat workbook SHA-256 mismatch")
    mapping = commune_mapping()
    records = []
    labels = {}
    value_labels: dict[str, dict[str, str]] = defaultdict(dict)
    workbook = load_workbook(RAW_PATH, read_only=True, data_only=True)
    try:
        for sheet, expected_headers in SHEETS.items():
            source_rows = iter(workbook[sheet].iter_rows(values_only=True))
            headers = tuple(next(source_rows))
            if headers != expected_headers:
                raise ValueError(f"Header drift in {sheet}: {headers!r}")
            dimensions_end = VALUE_COLUMNS[sheet][0]
            groupby_header = headers[1] if sheet != "Jaaruitgave" else "Stelsel"
            rows = []
            seen = set()
            for row_number, values in enumerate(source_rows, 2):
                if values[0] != 2025:
                    continue
                native = dict(zip(headers[:dimensions_end], values[:dimensions_end]))
                semantic = json.dumps(native, sort_keys=True, ensure_ascii=False)
                value_id = hashlib.sha256(semantic.encode()).hexdigest()[:24]
                if value_id in seen:
                    raise ValueError(f"Duplicate native category in {sheet}: {native}")
                seen.add(value_id)
                filters = {}
                guards = []
                for index, header in enumerate(headers[:dimensions_end]):
                    value = values[index]
                    guards.append(
                        {
                            "column": get_column_letter(index + 1),
                            "expected_value": value,
                            "label": header,
                        }
                    )
                    if index == 0 or value == "":
                        continue
                    key = dimension(header)
                    filters[key] = value
                    labels[key] = header
                    value_labels[key][str(value)] = str(value)
                for index in VALUE_COLUMNS[sheet]:
                    if isinstance(values[index], bool) or not isinstance(
                        values[index], (int, float)
                    ):
                        raise ValueError(
                            f"Non-numeric published cell {sheet}!{get_column_letter(index + 1)}{row_number}"
                        )
                    if (
                        headers[index] == "Aantal"
                        and int(values[index]) != values[index]
                    ):
                        raise ValueError(
                            f"Non-integer published count in {sheet} row {row_number}"
                        )
                rows.append(
                    {
                        "value_id": value_id,
                        "label": " / ".join(str(v) for v in values[1:dimensions_end]),
                        "ordinal": len(rows),
                        "row_number": row_number,
                        "source_row_id": value_id,
                        "expected_row_header_column": "A",
                        "expected_row_header": 2025,
                        "table_record_kind": "total"
                        if sheet == "Gepensioneerden_gemeente"
                        else "detail",
                        **geography(native, mapping),
                        "filters": filters,
                        "constraints": [
                            {
                                "variable": key,
                                "operator": "==",
                                "value": value,
                                "label": labels[key],
                            }
                            for key, value in filters.items()
                        ],
                        "guard_cells": guards,
                    }
                )
            annual = sheet == "Jaaruitgave"
            base_id = "pensionstat.2025." + sheet.lower()
            groups = (
                IGO_RECORD_SETS
                if sheet == "IGO"
                else ((None, VALUE_COLUMNS[sheet], "annual" if annual else "january"),)
            )
            # Split sets share the sheet's spec ID and record-ID prefix, so
            # fact keys and source-record IDs do not depend on the split.
            for suffix, columns, coverage in groups:
                records.append(
                    {
                        "record_set_id": f"{base_id}.{suffix}" if suffix else base_id,
                        "provenance_class": "administrative",
                        "assertion": "observation",
                        "record_set_spec_id": "pensionstat." + sheet.lower() + ".v1",
                        "source_record_id_prefix": base_id,
                        "sheet_name": sheet,
                        "period_type": "calendar_year" if annual else "month",
                        "period": 2025 if annual else "2025-01",
                        "period_coverage": COVERAGE[coverage],
                        "geography_id": "BE",
                        "geography_level": "country",
                        "geography_name": "Belgium",
                        "geography_vintage": "current",
                        "entity": "person",
                        "entity_role": "benefit_recipient",
                        "domain": "pensions",
                        "groupby_dimension": dimension(groupby_header),
                        "rows": rows,
                        "measures": [
                            measure(sheet, headers[index], index, ordinal)
                            for ordinal, index in enumerate(columns)
                        ],
                    }
                )
    finally:
        workbook.close()
    return {
        "schema_version": "ledger.source_package.v1",
        "package_id": PACKAGE_ID,
        "label": "PensionStat 2025 legal pensioners, IGO, expenditure and minimum pensions",
        "dimension_labels": labels,
        "dimension_value_labels": dict(value_labels),
        "artifact": {
            "source_name": "sfpd_pensions",
            "source_table": SOURCE_TABLE,
            "resource_package": "db",
            "resource_directory": "data/sfpd/legal_pension_caseload_2025",
            "manifest": "manifest.yaml",
            "vintage": "pensionstat_2025_workbook",
            "extracted_at": "2026-10-23",
            "extraction_method": "Direct cell selection from verbatim PensionStat workbook; five complete worksheets parsed",
            "parser": "xlsx_read_only_used_range",
            "artifact_year": 2025,
            "sheets": ["Info", *SHEETS],
        },
        "record_sets": records,
    }


def main() -> None:
    package = generate_package()
    (PACKAGE_DIR / "source_package.yaml").write_text(
        yaml_text(package), encoding="utf-8"
    )
    mapping = commune_mapping()
    communes = next(
        r
        for r in package["record_sets"]
        if r["sheet_name"] == "Gepensioneerden_gemeente"
    )
    audit = []
    for row in communes["rows"]:
        name = row["filters"]["pensionstat.gemeente"]
        audit.append(
            {
                "source_sheet": communes["sheet_name"],
                "source_cell": f"C{row['row_number']}",
                "publisher_name": name,
                **mapping[name],
            }
        )
    (PACKAGE_DIR / "commune_mapping.json").write_text(
        json.dumps(audit, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    print(
        json.dumps(
            {
                r["record_set_id"]: len(r["rows"]) * len(r["measures"])
                for r in package["record_sets"]
            }
        )
    )


if __name__ == "__main__":
    main()
