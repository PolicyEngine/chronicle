"""Source-row records for preserving full delimited source artifacts."""

from __future__ import annotations

import csv
import hashlib
import itertools
import json
import re
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import date, datetime
from io import BytesIO, StringIO
from itertools import product
from pathlib import Path
from typing import Any

import openpyxl

from chronicle.epoch import EMIT_EPOCH, Epoch, canonicalize_key, hash_domain
from chronicle.sources.cells import (
    SourceArtifactMetadata,
    SourceCell,
    decode_delimited_text,
)

Scalar = str | int | float | bool | None
SOURCE_ROW_KEY_PREFIX = hash_domain("source_row")
SOURCE_COLUMN_KEY_PREFIX = hash_domain("source_column")
SOURCE_ROW_VALUE_KEY_PREFIX = hash_domain("source_row_value")


@dataclass(frozen=True)
class SourceRow:
    """One parsed row from a source artifact."""

    artifact: SourceArtifactMetadata
    sheet_name: str
    row_number: int
    values: dict[str, Scalar]


@dataclass(frozen=True)
class SourceColumn:
    """One parsed column from a row-oriented source artifact."""

    artifact: SourceArtifactMetadata
    sheet_name: str
    column_number: int
    raw_name: str
    normalized_name: str


@dataclass(frozen=True)
class SourceRowValue:
    """One queryable source-row value at a row/column coordinate."""

    source_row_key: str
    source_column_key: str
    row_number: int
    column_number: int
    raw_column_name: str
    normalized_column_name: str
    value: Scalar


@dataclass(frozen=True)
class SourceRowIssue:
    """One source-row validation issue."""

    code: str
    message: str
    source_row_key: str | None = None
    row_index: int | None = None

    def to_dict(self) -> dict[str, Any]:
        """Return a JSON-serializable issue."""
        data = asdict(self)
        return {key: value for key, value in data.items() if value is not None}


@dataclass(frozen=True)
class SourceRowReport:
    """Validation and QA summary for source-row records."""

    row_count: int
    counts: dict[str, dict[str, int]]
    errors: tuple[SourceRowIssue, ...]
    warnings: tuple[SourceRowIssue, ...] = ()

    @property
    def valid(self) -> bool:
        """Whether the row set has no validation errors."""
        return not self.errors

    def to_dict(self) -> dict[str, Any]:
        """Return a JSON-serializable report."""
        return {
            "valid": self.valid,
            "row_count": self.row_count,
            "counts": self.counts,
            "errors": [issue.to_dict() for issue in self.errors],
            "warnings": [issue.to_dict() for issue in self.warnings],
        }


def build_source_row_key(
    row: SourceRow,
    *,
    epoch: Epoch = EMIT_EPOCH,
) -> str:
    """Build a stable key from artifact hash and row coordinate."""
    payload = {
        "artifact_sha256": row.artifact.sha256,
        "sheet_name": row.sheet_name,
        "row_number": row.row_number,
    }
    raw = json.dumps(payload, sort_keys=True, separators=(",", ":"))
    digest = hashlib.sha256(raw.encode("utf-8")).hexdigest()[:24]
    return f"{hash_domain('source_row', epoch)}:{digest}"


def build_source_column_key(
    column: SourceColumn,
    *,
    epoch: Epoch = EMIT_EPOCH,
) -> str:
    """Build a stable key from artifact hash and column coordinate."""
    payload = {
        "artifact_sha256": column.artifact.sha256,
        "sheet_name": column.sheet_name,
        "column_number": column.column_number,
    }
    raw = json.dumps(payload, sort_keys=True, separators=(",", ":"))
    digest = hashlib.sha256(raw.encode("utf-8")).hexdigest()[:24]
    return f"{hash_domain('source_column', epoch)}:{digest}"


def build_source_row_value_key(
    row_value: SourceRowValue,
    *,
    epoch: Epoch = EMIT_EPOCH,
) -> str:
    """Build a stable key from source row and column keys."""
    payload = {
        "source_row_key": canonicalize_key("source_row", row_value.source_row_key),
        "source_column_key": canonicalize_key(
            "source_column", row_value.source_column_key
        ),
    }
    raw = json.dumps(payload, sort_keys=True, separators=(",", ":"))
    digest = hashlib.sha256(raw.encode("utf-8")).hexdigest()[:24]
    return f"{hash_domain('source_row_value', epoch)}:{digest}"


def source_rows_from_delimited_text(
    content: bytes,
    artifact: SourceArtifactMetadata,
    *,
    sheet_name: str,
    delimiter: str = ",",
    header_row: int = 1,
) -> list[SourceRow]:
    """Parse every data row from a delimited text artifact.

    ``header_row`` names the physical line carrying the column headers.
    Export tools such as SuperWEB2 write a metadata preamble above the table;
    lines before the header are preamble, not data, and are skipped. Row
    numbers stay physical file line numbers either way. Empty content yields
    no rows, but content that ends before the declared header line is a
    package error and refuses loudly rather than emitting nothing.
    """
    if header_row < 1:
        raise ValueError(
            f"header_row is a 1-based physical line number, got {header_row}"
        )
    text = decode_delimited_text(content)
    reader = csv.reader(StringIO(text), delimiter=delimiter)
    header: list[str] | None = None
    rows: list[SourceRow] = []
    last_line_number = 0
    for source_line_number, raw_row in enumerate(reader, start=1):
        last_line_number = source_line_number
        if source_line_number < header_row:
            continue
        if source_line_number == header_row:
            header = [_normalize_header_cell(item) for item in raw_row]
            continue
        values = {
            column: _delimited_scalar(raw_row[index]) if index < len(raw_row) else None
            for index, column in enumerate(header or ())
        }
        rows.append(
            SourceRow(
                artifact=artifact,
                sheet_name=sheet_name,
                row_number=source_line_number,
                values=values,
            )
        )
    if header is None and last_line_number:
        raise ValueError(
            f"Delimited text ends at line {last_line_number} before the "
            f"declared header row {header_row}"
        )
    return rows


def source_rows_from_xlsx_table(
    content: bytes,
    artifact: SourceArtifactMetadata,
    *,
    sheet_name: str,
    header_row: int = 1,
) -> list[SourceRow]:
    """Parse one worksheet's rectangular table into full source rows.

    The delimited full-row contract applied to a workbook sheet: the whole
    table is preserved as source rows keyed by the header line's column names,
    and packages select rows at the record-set or artifact level. This is the
    lane for wide publisher workbooks whose used-range cell parse would cost
    millions of cell records per build (the PIPR monthly workbook is one
    ~49k-row sheet); the artifact's ``selected_rows`` then restricts cell
    emission while every row stays preserved and queryable.

    Row numbers are physical sheet rows. Datetime cells flatten to ISO text
    (date-only at midnight) so criteria and guards compare against a stable
    string form. A sheet that ends before the declared header row refuses
    loudly — a package silently contributing nothing is invisible until
    someone diffs fact counts.
    """
    if header_row < 1:
        raise ValueError(
            f"header_row is a 1-based physical sheet row, got {header_row}"
        )
    workbook = openpyxl.load_workbook(
        BytesIO(content),
        read_only=True,
        data_only=True,
    )
    # Read-only mode holds its source open until close() is called; this lane
    # exists for memory pressure, so release the handle deterministically.
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Workbook does not carry the declared sheet {sheet_name!r}"
            )
        sheet = workbook[sheet_name]
        header: list[str] | None = None
        rows: list[SourceRow] = []
        last_row_number = 0
        for row_number, raw_row in enumerate(
            sheet.iter_rows(values_only=True),
            start=1,
        ):
            last_row_number = row_number
            if row_number < header_row:
                continue
            if row_number == header_row:
                header = [
                    _normalize_header_cell("" if item is None else str(item))
                    for item in raw_row
                ]
                continue
            values = {
                column: _xlsx_row_scalar(raw_row[index])
                if index < len(raw_row)
                else None
                for index, column in enumerate(header or ())
            }
            rows.append(
                SourceRow(
                    artifact=artifact,
                    sheet_name=sheet_name,
                    row_number=row_number,
                    values=values,
                )
            )
        if header is None:
            raise ValueError(
                f"Sheet {sheet_name!r} ends at row {last_row_number} before "
                f"the declared header row {header_row}"
            )
        return rows
    finally:
        workbook.close()


def _xlsx_row_scalar(value: Any) -> Scalar:
    """Normalize one xlsx cell value into a source-row scalar."""
    if isinstance(value, datetime):
        if (value.hour, value.minute, value.second, value.microsecond) == (0, 0, 0, 0):
            return value.date().isoformat()
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        return value.strip()
    return value


def source_rows_from_json_table(
    content: bytes,
    artifact: SourceArtifactMetadata,
    *,
    sheet_name: str,
) -> list[SourceRow]:
    """Parse a simple JSON table into source rows."""
    data = json.loads(content.decode("utf-8"))
    if not isinstance(data, list) or not data:
        return []

    first = data[0]
    if isinstance(first, list):
        header = [_normalize_header_cell(str(item)) for item in first]
        rows: list[SourceRow] = []
        for source_row_number, raw_row in enumerate(data[1:], start=2):
            if not isinstance(raw_row, list):
                raise ValueError("JSON table rows must all be arrays.")
            rows.append(
                SourceRow(
                    artifact=artifact,
                    sheet_name=sheet_name,
                    row_number=source_row_number,
                    values={
                        column: _json_scalar(raw_row[index])
                        if index < len(raw_row)
                        else None
                        for index, column in enumerate(header)
                    },
                )
            )
        return rows

    if isinstance(first, dict):
        header = list(first)
        rows = []
        for source_row_number, raw_row in enumerate(data, start=1):
            if not isinstance(raw_row, dict):
                raise ValueError("JSON table rows must all be objects.")
            rows.append(
                SourceRow(
                    artifact=artifact,
                    sheet_name=sheet_name,
                    row_number=source_row_number,
                    values={
                        _normalize_header_cell(str(column)): _json_scalar(
                            raw_row.get(column)
                        )
                        for column in header
                    },
                )
            )
        return rows

    raise ValueError("JSON table must be an array of arrays or objects.")


def source_rows_from_json_stat_2(
    content: bytes,
    artifact: SourceArtifactMetadata,
    *,
    sheet_name: str,
) -> list[SourceRow]:
    """Flatten a JSON-stat 2.0 dataset cube into deterministic source rows.

    JSON-stat stores observations in row-major order, with the last dimension
    changing fastest. The returned rows preserve that order, retain every cube
    position (including null observations), and expose dimension codes before
    the observation value so declarative packages can select exact series.
    """
    data = json.loads(content.decode("utf-8"))
    if not isinstance(data, dict):
        raise ValueError("JSON-stat 2.0 payload must be an object.")
    if data.get("class") != "dataset":
        raise ValueError("JSON-stat 2.0 payload class must be 'dataset'.")
    if data.get("version") != "2.0":
        raise ValueError("JSON-stat payload version must be '2.0'.")

    dimension_ids = data.get("id")
    sizes = data.get("size")
    dimensions = data.get("dimension")
    if not isinstance(dimension_ids, list) or not dimension_ids:
        raise ValueError("JSON-stat dataset id must be a non-empty array.")
    if not all(isinstance(item, str) and item for item in dimension_ids):
        raise ValueError("JSON-stat dataset dimension IDs must be non-empty strings.")
    if len(set(dimension_ids)) != len(dimension_ids):
        raise ValueError("JSON-stat dataset dimension IDs must be unique.")
    if not isinstance(sizes, list) or len(sizes) != len(dimension_ids):
        raise ValueError("JSON-stat dataset size must align with id.")
    if not all(type(size) is int and size > 0 for size in sizes):
        raise ValueError("JSON-stat dataset sizes must be positive integers.")
    if not isinstance(dimensions, dict):
        raise ValueError("JSON-stat dataset dimension must be an object.")

    codes_by_dimension: list[list[str]] = []
    labels_by_dimension: list[dict[str, str]] = []
    for dimension_id, size in zip(dimension_ids, sizes, strict=True):
        codes, labels = _json_stat_dimension_categories(
            dimensions,
            dimension_id=dimension_id,
            expected_size=size,
        )
        codes_by_dimension.append(codes)
        labels_by_dimension.append(labels)

    observation_count = 1
    for size in sizes:
        observation_count *= size
    values = _json_stat_observations(
        data.get("value"),
        field="value",
        observation_count=observation_count,
        required=True,
    )
    statuses = _json_stat_observations(
        data.get("status"),
        field="status",
        observation_count=observation_count,
        required=False,
    )

    rows: list[SourceRow] = []
    coordinates = product(*(range(size) for size in sizes))
    for source_index, positions in enumerate(coordinates):
        dimension_codes = {
            dimension_id: codes_by_dimension[index][position]
            for index, (dimension_id, position) in enumerate(
                zip(dimension_ids, positions, strict=True)
            )
        }
        row_values: dict[str, Scalar] = dict(dimension_codes)
        row_values["value"] = _json_stat_scalar(values[source_index], field="value")
        row_values["status"] = _json_stat_scalar(
            statuses[source_index],
            field="status",
        )
        row_values["source_index"] = source_index
        for index, dimension_id in enumerate(dimension_ids):
            code = dimension_codes[dimension_id]
            row_values[f"{dimension_id}_label"] = labels_by_dimension[index].get(
                code,
                code,
            )
        rows.append(
            SourceRow(
                artifact=artifact,
                sheet_name=sheet_name,
                row_number=source_index + 1,
                values=row_values,
            )
        )
    return rows


def _json_stat_dimension_categories(
    dimensions: dict[str, Any],
    *,
    dimension_id: str,
    expected_size: int,
) -> tuple[list[str], dict[str, str]]:
    dimension = dimensions.get(dimension_id)
    if not isinstance(dimension, dict):
        raise ValueError(f"JSON-stat dimension {dimension_id!r} is missing.")
    category = dimension.get("category")
    if not isinstance(category, dict):
        raise ValueError(
            f"JSON-stat dimension {dimension_id!r} category must be an object."
        )
    category_index = category.get("index")
    if isinstance(category_index, list):
        if not all(isinstance(code, str) and code for code in category_index):
            raise ValueError(
                f"JSON-stat dimension {dimension_id!r} category index values "
                "must be non-empty strings."
            )
        codes = list(category_index)
    elif isinstance(category_index, dict):
        positions: dict[int, str] = {}
        for code, position in category_index.items():
            if not isinstance(code, str) or not code:
                raise ValueError(
                    f"JSON-stat dimension {dimension_id!r} category codes "
                    "must be non-empty strings."
                )
            if type(position) is not int or position < 0:
                raise ValueError(
                    f"JSON-stat dimension {dimension_id!r} category positions "
                    "must be non-negative integers."
                )
            if position in positions:
                raise ValueError(
                    f"JSON-stat dimension {dimension_id!r} category positions "
                    "must be unique."
                )
            positions[position] = code
        if set(positions) != set(range(expected_size)):
            raise ValueError(
                f"JSON-stat dimension {dimension_id!r} category positions must "
                f"cover 0 through {expected_size - 1}."
            )
        codes = [positions[position] for position in range(expected_size)]
    else:
        raise ValueError(
            f"JSON-stat dimension {dimension_id!r} category index must be an "
            "array or object."
        )

    if len(codes) != expected_size or len(set(codes)) != len(codes):
        raise ValueError(
            f"JSON-stat dimension {dimension_id!r} categories must contain "
            f"{expected_size} unique codes."
        )

    raw_labels = category.get("label", {})
    if not isinstance(raw_labels, dict):
        raise ValueError(
            f"JSON-stat dimension {dimension_id!r} category labels must be an object."
        )
    labels = {
        code: str(raw_labels[code])
        for code in codes
        if raw_labels.get(code) is not None
    }
    return codes, labels


def _json_stat_observations(
    observations: Any,
    *,
    field: str,
    observation_count: int,
    required: bool,
) -> list[Any]:
    if observations is None:
        if required:
            raise ValueError(f"JSON-stat dataset {field} is required.")
        return [None] * observation_count
    if isinstance(observations, list):
        if len(observations) != observation_count:
            raise ValueError(
                f"JSON-stat dataset {field} array must contain "
                f"{observation_count} entries."
            )
        return list(observations)
    if isinstance(observations, dict):
        dense = [None] * observation_count
        for raw_index, value in observations.items():
            try:
                index = int(raw_index)
            except (TypeError, ValueError) as exc:
                raise ValueError(
                    f"JSON-stat dataset {field} keys must be integer indexes."
                ) from exc
            if str(index) != str(raw_index) or not 0 <= index < observation_count:
                raise ValueError(
                    f"JSON-stat dataset {field} index {raw_index!r} is out of range."
                )
            dense[index] = value
        return dense
    raise ValueError(f"JSON-stat dataset {field} must be an array or object.")


def _json_stat_scalar(value: Any, *, field: str) -> Scalar:
    if value is None or isinstance(value, bool | int | float | str):
        return _json_scalar(value)
    raise ValueError(f"JSON-stat dataset {field} entries must be scalar values.")


S0101_TOTAL_AGE_COLUMNS: tuple[tuple[str, str], ...] = (
    ("S0101_C01_002E", "Aged 0-4"),
    ("S0101_C01_003E", "Aged 5-9"),
    ("S0101_C01_004E", "Aged 10-14"),
    ("S0101_C01_005E", "Aged 15-19"),
    ("S0101_C01_006E", "Aged 20-24"),
    ("S0101_C01_007E", "Aged 25-29"),
    ("S0101_C01_008E", "Aged 30-34"),
    ("S0101_C01_009E", "Aged 35-39"),
    ("S0101_C01_010E", "Aged 40-44"),
    ("S0101_C01_011E", "Aged 45-49"),
    ("S0101_C01_012E", "Aged 50-54"),
    ("S0101_C01_013E", "Aged 55-59"),
    ("S0101_C01_014E", "Aged 60-64"),
    ("S0101_C01_015E", "Aged 65-69"),
    ("S0101_C01_016E", "Aged 70-74"),
    ("S0101_C01_017E", "Aged 75-79"),
    ("S0101_C01_018E", "Aged 80-84"),
    ("S0101_C01_019E", "Aged 85 and over"),
)


B01001_FEMALE_AGE_COLUMNS: tuple[tuple[str, str, int, int], ...] = (
    ("B01001_030E", "Female 15 to 17 years", 15, 18),
    ("B01001_031E", "Female 18 and 19 years", 18, 20),
    ("B01001_032E", "Female 20 years", 20, 21),
    ("B01001_033E", "Female 21 years", 21, 22),
    ("B01001_034E", "Female 22 to 24 years", 22, 25),
    ("B01001_035E", "Female 25 to 29 years", 25, 30),
    ("B01001_036E", "Female 30 to 34 years", 30, 35),
    ("B01001_037E", "Female 35 to 39 years", 35, 40),
    ("B01001_038E", "Female 40 to 44 years", 40, 45),
)


S2201_SNAP_HOUSEHOLD_COLUMNS: tuple[tuple[str, str, str], ...] = (
    (
        "S2201_C01_001E",
        "all",
        "Estimate!!Total!!Households",
    ),
    (
        "S2201_C03_001E",
        "receiving_food_stamps_snap",
        "Estimate!!Households receiving food stamps/SNAP!!Households",
    ),
    (
        "S2201_C05_001E",
        "not_receiving_food_stamps_snap",
        "Estimate!!Households not receiving food stamps/SNAP!!Households",
    ),
)


def source_rows_from_census_acs_s0101_age_json(
    content: bytes,
    artifact: SourceArtifactMetadata,
    *,
    sheet_name: str,
) -> list[SourceRow]:
    """Unpivot Census ACS S0101 API rows into age-band source rows."""
    table_rows = source_rows_from_json_table(
        content,
        artifact,
        sheet_name=sheet_name,
    )
    rows: list[SourceRow] = []
    for source_table_row in table_rows:
        for variable, age_label in S0101_TOTAL_AGE_COLUMNS:
            rows.append(
                SourceRow(
                    artifact=artifact,
                    sheet_name=sheet_name,
                    row_number=len(rows) + 1,
                    values={
                        "GEO_ID": source_table_row.values.get("GEO_ID"),
                        "NAME": source_table_row.values.get("NAME"),
                        "source_column_id": variable,
                        "source_concept": age_label,
                        "age": age_label,
                        "value": source_table_row.values.get(variable),
                        "source_table_row_number": source_table_row.row_number,
                    },
                )
            )
    return rows


def source_rows_from_census_acs_s2201_snap_json(
    content: bytes,
    artifact: SourceArtifactMetadata,
    *,
    sheet_name: str,
) -> list[SourceRow]:
    """Unpivot Census ACS S2201 API rows into SNAP household-count rows."""
    table_rows = source_rows_from_json_table(
        content,
        artifact,
        sheet_name=sheet_name,
    )
    rows: list[SourceRow] = []
    for source_table_row in table_rows:
        for variable, snap_status, source_concept in S2201_SNAP_HOUSEHOLD_COLUMNS:
            rows.append(
                SourceRow(
                    artifact=artifact,
                    sheet_name=sheet_name,
                    row_number=len(rows) + 1,
                    values={
                        "GEO_ID": source_table_row.values.get("GEO_ID"),
                        "NAME": source_table_row.values.get("NAME"),
                        "snap_receipt_status": snap_status,
                        "source_column_id": variable,
                        "source_concept": source_concept,
                        "value": source_table_row.values.get(variable),
                        "source_table_row_number": source_table_row.row_number,
                    },
                )
            )
    return rows


def source_rows_from_census_b01001_female_age_json(
    content: bytes,
    artifact: SourceArtifactMetadata,
    *,
    sheet_name: str,
) -> list[SourceRow]:
    """Unpivot Census ACS B01001 API rows into female age-band source rows."""
    table_rows = source_rows_from_json_table(
        content,
        artifact,
        sheet_name=sheet_name,
    )
    rows: list[SourceRow] = []
    for source_table_row in table_rows:
        state_value = source_table_row.values.get("state")
        if state_value is None:
            raise ValueError("Census B01001 API row is missing state FIPS.")
        state_code = f"{int(state_value):02d}"
        for variable, age_label, _lower_age, _upper_age in B01001_FEMALE_AGE_COLUMNS:
            rows.append(
                SourceRow(
                    artifact=artifact,
                    sheet_name=sheet_name,
                    row_number=len(rows) + 1,
                    values={
                        "state": state_code,
                        "geography_id": f"0400000US{state_code}",
                        "sex": "female",
                        "source_column_id": variable,
                        "source_concept": age_label,
                        "age": age_label,
                        "value": source_table_row.values.get(variable),
                        "source_table_row_number": source_table_row.row_number,
                    },
                )
            )
    return rows


MONTH_NAME_TO_NUMBER: dict[str, int] = {
    "January": 1,
    "February": 2,
    "March": 3,
    "April": 4,
    "May": 5,
    "June": 6,
    "July": 7,
    "August": 8,
    "September": 9,
    "October": 10,
    "November": 11,
    "December": 12,
}


def source_rows_from_cdc_vsrr_live_births_json(
    content: bytes,
    artifact: SourceArtifactMetadata,
    *,
    sheet_name: str,
) -> list[SourceRow]:
    """Parse CDC VSRR live-birth Socrata rows with month period values."""
    data = json.loads(content.decode("utf-8"))
    if not isinstance(data, list):
        raise ValueError("CDC VSRR JSON response must be an array of objects.")

    rows: list[SourceRow] = []
    for source_row_number, raw_row in enumerate(data, start=1):
        if not isinstance(raw_row, dict):
            raise ValueError("CDC VSRR JSON rows must all be objects.")
        month = str(raw_row.get("month") or "").strip()
        month_number = MONTH_NAME_TO_NUMBER.get(month)
        if month_number is None:
            raise ValueError(f"Unsupported CDC VSRR month label: {month!r}")
        year = int(_json_scalar(raw_row.get("year")))
        rows.append(
            SourceRow(
                artifact=artifact,
                sheet_name=sheet_name,
                row_number=source_row_number,
                values={
                    "state": _json_scalar(raw_row.get("state")),
                    "year": year,
                    "month": month,
                    "month_number": month_number,
                    "period": f"{year}-{month_number:02d}",
                    "frequency": _json_scalar(raw_row.get("period")),
                    "indicator": _json_scalar(raw_row.get("indicator")),
                    "data_value": _json_scalar(raw_row.get("data_value")),
                },
            )
        )
    return rows


def source_rows_from_ons_timeseries_json(
    content: bytes,
    artifact: SourceArtifactMetadata,
    *,
    sheet_name: str = "years",
) -> list[SourceRow]:
    """Parse ONS time-series JSON observations into source rows."""
    data = json.loads(content.decode("utf-8"))
    description = data.get("description") or {}
    if not isinstance(description, dict):
        description = {}
    rows: list[SourceRow] = []
    row_number = 1
    for source_field, frequency in (
        ("years", "annual"),
        ("quarters", "quarterly"),
        ("months", "monthly"),
    ):
        for item in data.get(source_field) or []:
            if not isinstance(item, dict):
                continue
            year = _delimited_scalar(str(item.get("year") or ""))
            rows.append(
                SourceRow(
                    artifact=artifact,
                    sheet_name=sheet_name,
                    row_number=row_number,
                    values={
                        "ons_series_id": description.get("cdid"),
                        "dataset_id": description.get("datasetId"),
                        "frequency": frequency,
                        "period": year if frequency == "annual" else item.get("date"),
                        "year": year,
                        "quarter": item.get("quarter") or None,
                        "month": item.get("month") or None,
                        "value": _delimited_scalar(str(item.get("value") or "")),
                        "date": item.get("date"),
                        "label": item.get("label"),
                        "source_dataset": item.get("sourceDataset"),
                        "update_date": item.get("updateDate"),
                        "release_date": description.get("releaseDate"),
                        "latest_period": description.get("date"),
                        "title": description.get("title"),
                    },
                )
            )
            row_number += 1
    return rows


def source_rows_from_ees_permalink_table_html(
    content: bytes,
    artifact: SourceArtifactMetadata,
    *,
    sheet_name: str = "table",
) -> list[SourceRow]:
    """Parse an Explore Education Statistics permalink table into source rows."""
    text = content.decode("utf-8")
    match = re.search(
        r'<script id="__NEXT_DATA__" type="application/json">(.+?)</script>',
        text,
        flags=re.DOTALL,
    )
    if not match:
        raise ValueError("Could not find __NEXT_DATA__ table payload")

    next_data = json.loads(match.group(1))
    table_json = next_data["props"]["pageProps"]["data"]["table"]["json"]
    column_groups = _ees_value_column_groups(table_json)
    value_columns = _ees_year_value_columns(column_groups)
    total_value_cell_count = sum(len(group["years"]) for group in column_groups)

    rows: list[SourceRow] = []
    current_row_group: str | None = None
    virtual_row_number = 1
    for row_number, raw_row in enumerate(table_json.get("tbody") or [], start=1):
        cells = [_ees_cell_text(cell) for cell in raw_row]
        if not cells:
            continue

        if len(cells) == total_value_cell_count + 2:
            row_group = cells[0]
            row_measure = cells[1]
            value_cells = cells[2:]
            if row_group:
                current_row_group = row_group
        elif len(cells) == total_value_cell_count + 1:
            if current_row_group is None:
                raise ValueError(
                    "EES continuation row appeared before any row group; "
                    f"source row {row_number}."
                )
            row_group = current_row_group
            row_measure = cells[0]
            value_cells = cells[1:]
        else:
            raise ValueError(
                "EES table row has an unexpected cell count; "
                f"source row {row_number} has {len(cells)} cells, expected "
                f"{total_value_cell_count + 2} grouped or "
                f"{total_value_cell_count + 1} continuation cells."
            )

        offset = 0
        for group in column_groups:
            values: dict[str, Scalar] = {
                "row_group": row_group,
                "row_measure": row_measure,
                "repayment_plan": group["repayment_plan"],
                "borrower_status": _ees_borrower_status(row_measure),
                "source_table_row_number": row_number,
            }
            group_values = value_cells[offset : offset + len(group["years"])]
            offset += len(group["years"])
            values.update({column: None for column in value_columns})
            for year, raw_value in zip(group["years"], group_values):
                values[f"value_{year}"] = (
                    None if raw_value == "no data" else _delimited_scalar(raw_value)
                )
            rows.append(
                SourceRow(
                    artifact=artifact,
                    sheet_name=sheet_name,
                    row_number=virtual_row_number,
                    values=values,
                )
            )
            virtual_row_number += 1
        if offset != len(value_cells):
            raise ValueError(
                f"EES parser did not consume all value cells; source row {row_number}."
            )
    return rows


def source_rows_from_kff_state_indicator_gdocs_html(
    content: bytes,
    artifact: SourceArtifactMetadata,
    *,
    sheet_name: str,
) -> list[SourceRow]:
    """Parse KFF State Health Facts embedded gdocsObject table rows."""
    text = content.decode("utf-8", errors="replace")
    match = re.search(r'"gdocsObject":(\[.*?\]),"postBody"', text, flags=re.S)
    if not match:
        raise ValueError("KFF State Health Facts page is missing gdocsObject data")

    gdocs_object = json.loads(match.group(1))
    rows: list[SourceRow] = []
    source_row_number = 0
    for sheet in gdocs_object:
        if not isinstance(sheet, list) or len(sheet) != 2:
            continue
        year_label, table = sheet
        year_text = str(year_label)
        if not year_text.isdigit() or not isinstance(table, list) or len(table) < 3:
            continue
        value_column = str(table[0][1]).strip() if len(table[0]) > 1 else "Value"
        unit = str(table[1][1]).strip() if len(table[1]) > 1 else None
        for table_row_number, table_row in enumerate(table[2:], start=3):
            if not isinstance(table_row, list) or len(table_row) < 2:
                continue
            source_row_number += 1
            rows.append(
                SourceRow(
                    artifact=artifact,
                    sheet_name=sheet_name,
                    row_number=source_row_number,
                    values={
                        "Year": int(year_text),
                        "Geography": _delimited_scalar(str(table_row[0])),
                        value_column: _delimited_scalar(str(table_row[1])),
                        "Unit": _delimited_scalar(unit or ""),
                        "KFF table row": table_row_number,
                    },
                )
            )
    return rows


def source_columns_from_source_rows(rows: list[SourceRow]) -> list[SourceColumn]:
    """Derive queryable source columns from parsed source rows."""
    columns_by_key: dict[str, SourceColumn] = {}
    for row in rows:
        for column_number, raw_name in enumerate(row.values, start=1):
            column = SourceColumn(
                artifact=row.artifact,
                sheet_name=row.sheet_name,
                column_number=column_number,
                raw_name=raw_name,
                normalized_name=normalize_source_column_name(raw_name),
            )
            columns_by_key.setdefault(build_source_column_key(column), column)
    return sorted(
        columns_by_key.values(),
        key=lambda column: (
            column.artifact.sha256,
            column.sheet_name,
            column.column_number,
        ),
    )


def source_row_values_from_source_rows(
    rows: list[SourceRow],
) -> list[SourceRowValue]:
    """Derive queryable source-row values from parsed source rows."""
    values: list[SourceRowValue] = []
    for row in rows:
        source_row_key = build_source_row_key(row)
        for column_number, (raw_name, value) in enumerate(
            row.values.items(),
            start=1,
        ):
            column = SourceColumn(
                artifact=row.artifact,
                sheet_name=row.sheet_name,
                column_number=column_number,
                raw_name=raw_name,
                normalized_name=normalize_source_column_name(raw_name),
            )
            values.append(
                SourceRowValue(
                    source_row_key=source_row_key,
                    source_column_key=build_source_column_key(column),
                    row_number=row.row_number,
                    column_number=column_number,
                    raw_column_name=raw_name,
                    normalized_column_name=column.normalized_name,
                    value=value,
                )
            )
    return values


def source_cells_from_source_rows(
    rows: list[SourceRow],
    *,
    selected_rows: tuple[dict[str, str], ...],
) -> list[SourceCell]:
    """Build compact selected source cells from full source rows."""
    if not rows:
        return []

    artifact = rows[0].artifact
    sheet_name = rows[0].sheet_name
    columns = list(rows[0].values)
    cells = [
        _source_row_cell(
            artifact,
            sheet_name,
            row_number=1,
            column_number=index + 1,
            raw_value=column,
        )
        for index, column in enumerate(columns)
    ]
    selected_by_index: dict[int, SourceRow] = {}
    if selected_rows:
        indices_by_key_tuple: dict[
            tuple[str, ...],
            dict[tuple[str, ...], list[SourceRow]],
        ] = {}
        for selection_index, criteria in enumerate(selected_rows):
            key_tuple = tuple(criteria)
            row_index = indices_by_key_tuple.get(key_tuple)
            if row_index is None:
                row_index = _source_row_index(rows, key_tuple)
                indices_by_key_tuple[key_tuple] = row_index
            matches = row_index.get(tuple(criteria.values()), [])
            if len(matches) != 1:
                raise ValueError(
                    "Selected row criteria must match exactly one source row; "
                    f"criteria {criteria!r} matched {len(matches)} rows."
                )
            selected_by_index[selection_index] = matches[0]
    else:
        selected_by_index = dict(enumerate(rows))

    for virtual_row_number, selection_index in enumerate(
        sorted(selected_by_index),
        start=2,
    ):
        row = selected_by_index[selection_index]
        source_row_key = build_source_row_key(row)
        for column_index, column in enumerate(columns, start=1):
            cells.append(
                _source_row_cell(
                    artifact,
                    sheet_name,
                    row_number=virtual_row_number,
                    column_number=column_index,
                    raw_value=row.values.get(column),
                    note=f"source_line_number={row.row_number}",
                    source_row_key=source_row_key,
                )
            )
    return cells


def validate_source_rows(rows: list[SourceRow]) -> SourceRowReport:
    """Validate source rows and return QA counts plus issues."""
    errors: list[SourceRowIssue] = []
    key_indices: dict[str, list[int]] = {}

    for index, row in enumerate(rows):
        key = build_source_row_key(row)
        key_indices.setdefault(key, []).append(index)
        if not row.artifact.source_name.strip():
            errors.append(
                SourceRowIssue(
                    code="missing_source_name",
                    message="Row artifact is missing source_name",
                    source_row_key=key,
                    row_index=index,
                )
            )
        if not row.artifact.sha256.strip():
            errors.append(
                SourceRowIssue(
                    code="missing_artifact_sha256",
                    message="Row artifact is missing sha256",
                    source_row_key=key,
                    row_index=index,
                )
            )
        if row.row_number < 1:
            errors.append(
                SourceRowIssue(
                    code="malformed_coordinate",
                    message="Row number must be one-based",
                    source_row_key=key,
                    row_index=index,
                )
            )
        if not row.values:
            errors.append(
                SourceRowIssue(
                    code="missing_values",
                    message="Source row has no values",
                    source_row_key=key,
                    row_index=index,
                )
            )

    for key, indices in key_indices.items():
        if len(indices) > 1:
            errors.append(
                SourceRowIssue(
                    code="duplicate_source_row_key",
                    message=f"Duplicate source-row key appears at indices {indices}",
                    source_row_key=key,
                    row_index=indices[0],
                )
            )

    return SourceRowReport(
        row_count=len(rows),
        counts=source_row_counts(rows),
        errors=tuple(errors),
    )


def source_row_counts(rows: list[SourceRow]) -> dict[str, dict[str, int]]:
    """Count rows across QA dimensions."""
    return {
        "by_source": _counter_dict(row.artifact.source_name for row in rows),
        "by_sheet": _counter_dict(row.sheet_name for row in rows),
    }


def load_source_rows_jsonl(path: str | Path) -> list[SourceRow]:
    """Load source rows from JSON Lines."""
    row_path = Path(path)
    rows: list[SourceRow] = []
    with row_path.open() as file:
        for line_number, line in enumerate(file, start=1):
            if not line.strip():
                continue
            try:
                payload = json.loads(line)
            except json.JSONDecodeError as exc:
                raise ValueError(
                    f"Invalid JSON on line {line_number} of {row_path}"
                ) from exc
            rows.append(source_row_from_mapping(payload))
    return rows


def save_source_rows_jsonl(rows: list[SourceRow], path: str | Path) -> None:
    """Save source rows to JSON Lines."""
    row_path = Path(path)
    row_path.parent.mkdir(parents=True, exist_ok=True)
    with row_path.open("w") as file:
        for row in rows:
            file.write(json.dumps(source_row_to_mapping(row), sort_keys=True))
            file.write("\n")


def source_row_from_mapping(payload: dict[str, Any]) -> SourceRow:
    """Build a source row from a JSON-compatible mapping."""
    return SourceRow(
        artifact=SourceArtifactMetadata(**payload["artifact"]),
        sheet_name=payload["sheet_name"],
        row_number=payload["row_number"],
        values=dict(payload["values"]),
    )


def source_row_to_mapping(row: SourceRow) -> dict[str, Any]:
    """Convert a source row to a JSON-compatible mapping."""
    return asdict(row)


def normalize_source_column_name(name: str) -> str:
    """Normalize a source column name for cross-source SQL queries."""
    spaced = re.sub(r"([a-z0-9])([A-Z])", r"\1_\2", name.strip())
    normalized = re.sub(r"[^A-Za-z0-9]+", "_", spaced).strip("_").lower()
    return normalized


def _row_matches(row: SourceRow, criteria: dict[str, str]) -> bool:
    return all(
        _value_as_selector_text(row.values.get(key)) == value
        for key, value in criteria.items()
    )


def _source_row_index(
    rows: list[SourceRow],
    keys: tuple[str, ...],
) -> dict[tuple[str, ...], list[SourceRow]]:
    indexed: dict[tuple[str, ...], list[SourceRow]] = {}
    for row in rows:
        key = tuple(_value_as_selector_text(row.values.get(column)) for column in keys)
        indexed.setdefault(key, []).append(row)
    return indexed


def _source_row_cell(
    artifact: SourceArtifactMetadata,
    sheet_name: str,
    *,
    row_number: int,
    column_number: int,
    raw_value: Scalar,
    note: str | None = None,
    source_row_key: str | None = None,
) -> SourceCell:
    return SourceCell(
        artifact=artifact,
        sheet_name=sheet_name,
        row_number=row_number,
        column_number=column_number,
        address=f"{_excel_column_name(column_number)}{row_number}",
        cell_type=_scalar_cell_type(raw_value),
        raw_value=raw_value,
        display_value=None if raw_value is None else str(raw_value),
        note=note,
        source_row_key=source_row_key,
    )


def _delimited_scalar(value: str) -> Scalar:
    stripped = value.strip()
    if not stripped:
        return None
    if "_" in stripped:
        return stripped
    numeric = stripped.replace("$", "").replace(",", "")
    if numeric.lstrip("-").isdigit():
        return int(numeric)
    try:
        return float(numeric)
    except ValueError:
        return stripped


def _json_scalar(value: Any) -> Scalar:
    if value is None or isinstance(value, bool | int | float):
        return value
    return _delimited_scalar(str(value))


def _value_as_selector_text(value: Scalar) -> str:
    return "" if value is None else str(value)


def _normalize_header_cell(value: str) -> str:
    return value.strip().removeprefix("%")


def _ees_value_column_groups(table_json: dict[str, Any]) -> list[dict[str, Any]]:
    header_rows = table_json.get("thead") or []
    if len(header_rows) < 2:
        raise ValueError("EES table payload must include two header rows")

    groups: list[dict[str, Any]] = []
    year_index = 0
    year_cells = header_rows[1]
    for cell in header_rows[0]:
        text = _ees_cell_text(cell)
        colspan = int(cell.get("colSpan") or 1)
        if not text:
            continue
        years = [
            _academic_year_end(_ees_cell_text(year_cell))
            for year_cell in year_cells[year_index : year_index + colspan]
        ]
        year_index += colspan
        groups.append(
            {
                "repayment_plan": normalize_source_column_name(text),
                "years": years,
            }
        )

    if year_index != len(year_cells):
        raise ValueError("EES grouped headers do not align with year headers")
    return groups


def _ees_year_value_columns(groups: list[dict[str, Any]]) -> list[str]:
    return [
        f"value_{year}"
        for year in sorted(
            {year for group in groups for year in group["years"]},
            reverse=True,
        )
    ]


def _academic_year_end(value: str) -> int:
    match = re.fullmatch(r"(\d{4})[-/]\d{2}", value.strip())
    if not match:
        raise ValueError(f"Unsupported academic year header: {value!r}")
    return int(match.group(1)) + 1


def _ees_cell_text(cell: dict[str, Any]) -> str:
    return str(cell.get("text") or "").strip()


def _ees_borrower_status(row_measure: str) -> str:
    normalized = normalize_source_column_name(row_measure)
    if "earning_above_repayment_threshold" in normalized:
        return "above_repayment_threshold"
    if normalized == "number_of_borrowers_liable_to_make_repayments":
        return "liable_to_repay"
    return normalized


def _scalar_cell_type(value: Scalar) -> str:
    if value is None:
        return "empty"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int | float):
        return "number"
    return "text"


def _excel_column_name(column_number: int) -> str:
    name = ""
    while column_number:
        column_number, remainder = divmod(column_number - 1, 26)
        name = f"{chr(65 + remainder)}{name}"
    return name


def _counter_dict(values: Any) -> dict[str, int]:
    return dict(sorted(Counter(values).items()))


def _statxplore_column_key(label: str) -> str:
    """Build a stable snake_case column key from a Stat-Xplore field label."""
    key = re.sub(r"[^a-z0-9]+", "_", label.lower()).strip("_")
    return key or "field"


def _statxplore_field_keys(fields: list[dict[str, Any]]) -> list[str]:
    """Return each Stat-Xplore field's source-row column key, in response order."""
    keys: list[str] = []
    for field in fields:
        key = _statxplore_column_key(str(field.get("label") or field.get("uri")))
        while key in keys:
            key = f"{key}_"
        keys.append(key)
    return keys


def statxplore_field_labels(content: bytes) -> dict[str, str]:
    """Return the publisher label of each Stat-Xplore field by its column key."""
    fields = json.loads(content.decode("utf-8")).get("fields") or []
    return {
        key: str(field["label"])
        for key, field in zip(_statxplore_field_keys(fields), fields)
        if field.get("label")
    }


def source_rows_from_statxplore_table(
    content: bytes,
    artifact: SourceArtifactMetadata,
    *,
    sheet_name: str,
) -> list[SourceRow]:
    """Unpivot a Stat-Xplore Open Data API /table response into source rows.

    The response is self-describing: it echoes the submitted query, names the
    measure, and lists each dimension field with its items (labels plus value
    URIs, whose trailing segments carry publisher codes such as GSS geography
    identifiers). One source row is emitted per cube cell, columns are the
    measure label, the cell value, and a label/URI pair per dimension field
    in response order.
    """
    data = json.loads(content.decode("utf-8"))
    fields = data.get("fields") or []
    cubes = data.get("cubes") or {}
    if not fields or not cubes:
        return []
    measure_uri, cube = next(iter(cubes.items()))
    measure_label = measure_uri
    for measure in data.get("measures") or []:
        if measure.get("uri") == measure_uri:
            measure_label = str(measure.get("label") or measure_uri)
            break
    field_keys = _statxplore_field_keys(fields)
    item_axes = [field.get("items") or [] for field in fields]
    values = cube.get("values") if isinstance(cube, dict) else None
    if values is None or any(not axis for axis in item_axes):
        return []
    rows: list[SourceRow] = []
    for indices in itertools.product(*(range(len(axis)) for axis in item_axes)):
        node: Any = values
        for index in indices:
            node = node[index]
        row_values: dict[str, Scalar] = {
            "measure": measure_label,
            "value": _json_scalar(node),
        }
        for key, axis, index in zip(field_keys, item_axes, indices):
            item = axis[index]
            labels = item.get("labels") or []
            uris = item.get("uris") or []
            row_values[key] = str(labels[0]) if labels else None
            row_values[f"{key}_uri"] = str(uris[0]) if uris else None
        rows.append(
            SourceRow(
                artifact=artifact,
                sheet_name=sheet_name,
                row_number=len(rows) + 1,
                values=row_values,
            )
        )
    return rows
